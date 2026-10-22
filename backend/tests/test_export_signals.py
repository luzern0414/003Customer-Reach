"""导出/采购历史/时效回归；全部公司与公告为隔离夹具，不是生产买家事实。"""
import uuid
from io import BytesIO
from openpyxl import load_workbook
from app.db import Session,Lead
from app.schemas import Candidate
from app.gateway import archive_sources
from app.service import save_candidate,ServiceError
import pytest


def fixture(client):
    """client为隔离HTTP客户端；建立有基础角色与独立公开公告的夹具，返回保存参数及ID，无外网。"""
    name='隔离采购回归-'+uuid.uuid4().hex[:8]
    conv=client.post('/api/conversations').json()['id']
    job=client.post('/api/conversations/'+conv+'/messages',json={'content':'你好','client_token':uuid.uuid4().hex}).json()['id']
    body=name+'是深圳灯具零售门店，主营吊灯。'
    quote='采购人：'+name+'。吊灯采购，提交截止：2099-12-31 18:00:00。'
    source=archive_sources(job,[{'url':'https://example.org/'+name,'title':'隔离资料','body':body}],'隔离测试')[0]
    purchase=archive_sources(job,[{'url':'https://example.org/'+name+'/bid','title':'隔离采购公告','body':quote}],'隔离测试','procurement')[0]
    candidate=Candidate(display_name=name,city='深圳',buyer_role='retailer',matched_products=['吊灯'],match_reason='隔离验证',facts=[
        {'field':'name','value':name,'quote':body,'source_id':source['id']},
        {'field':'buyer_role','value':'零售门店','quote':body,'source_id':source['id']}],
        procurement_signal={'source_id':purchase['id'],'buyer':name,'participant_role':'buyer','stage':'open','deadline':'2099-12-31 18:00:00','quote':quote})
    rows={source['id']:source,purchase['id']:purchase}
    saved=save_candidate(candidate,rows,job)
    return saved['id'],candidate,rows,job


def test_basic_enrichment_does_not_erase_purchase_history(client):
    """client为隔离客户端；无采购字段的补查保留已校验公告及其来源/截止日期。"""
    identifier,candidate,rows,job=fixture(client)
    save_candidate(candidate.model_copy(update={'procurement_signal':None}),rows,job)
    row=client.get('/api/leads/'+identifier).json()
    assert row['qualification']['signal_score']==10
    assert len(row['procurement_signals'])==1 and row['procurement_signals'][0]['display_status']=='公开采购'
    assert row['procurement_signals'][0]['deadline']=='2099-12-31 18:00:00'
    assert row['procurement_signals'][0]['url'].endswith('/bid')


def test_new_award_preserves_previous_open_record(client):
    """client为隔离客户端；中标公告不当当前采购，既有开放期原文仍保留。"""
    identifier,candidate,rows,job=fixture(client)
    text='采购人：'+candidate.display_name+'，吊灯采购中标结果。'
    source=archive_sources(job,[{'url':'https://example.org/'+candidate.display_name+'/award','title':'中标结果','body':text}],'隔离测试','procurement')[0]
    rows[source['id']]=source
    updated=candidate.model_copy(update={'procurement_signal':candidate.procurement_signal.model_copy(update={'source_id':source['id'],'quote':text,'stage':'awarded','deadline':''})})
    save_candidate(updated,rows,job)
    row=client.get('/api/leads/'+identifier).json()
    assert row['qualification']['signal_score']==0
    assert {x['stage'] for x in row['procurement_signals']}=={'open','awarded'}
    assert row['procurement_signals'][0]['display_status']=='中标历史'
    assert row['procurement_signals'][1]['display_status']=='历史公告'


def test_expired_signal_loses_score_before_list_sort_detail_and_export(client,monkeypatch):
    """client为隔离客户端，monkeypatch仅推进业务时间；过期不继续加分，原始stage不改。"""
    from app import service
    identifier,candidate,rows,job=fixture(client)
    before=client.get('/api/leads/'+identifier).json()
    monkeypatch.setattr(service,'now',lambda:'2100-01-01 00:00:00')
    listed=client.get('/api/leads?q='+candidate.display_name).json()['items'][0]
    detail=client.get('/api/leads/'+identifier).json()
    assert listed['score']==before['score']-10 and detail['score']==listed['score']
    assert detail['procurement_signals'][0]['stage']=='open' and detail['procurement_signals'][0]['display_status']=='已过期'
    workbook=load_workbook(BytesIO(client.get('/api/leads/export?q='+candidate.display_name).content))
    assert workbook['客户线索']['I2'].value==listed['score']
    assert '已过期' in workbook['客户线索']['P2'].value


def test_export_retains_alternatives_and_announcement_metadata(client):
    """client为隔离客户端；三张表保留替代值口径与公告阶段/日期，并标争议、不执行公式。"""
    identifier,candidate,rows,job=fixture(client)
    with Session.begin() as db:
        row=db.get(Lead,identifier)
        row.facts={**row.facts,'revenue':{'value':'100万元','verification':'disputed','scope':'母公司','effective_date':'2025','source_id':'first',
            'alternatives':[{'value':'200万元','scope':'门店','effective_date':'2024','source_id':'alternative','verification':'原文匹配'}]}}
    book=load_workbook(BytesIO(client.get('/api/leads/export?q='+candidate.display_name).content))
    assert book.sheetnames==['客户线索','事实与口径','逐字段证据']
    main=list(book['客户线索'].iter_rows(min_row=2,values_only=True))[0]
    headers=[cell.value for cell in book['客户线索'][1]]
    assert '年营收：争议' in main[headers.index('数据状态')] and main[headers.index('买家角色')]=='retailer'
    facts=list(book['事实与口径'].iter_rows(min_row=2,values_only=True))
    alternative=next(x for x in facts if x[8]=='替代值')
    assert alternative[3]=='200万元' and alternative[5:8]==('门店','2024','alternative')
    purchase=next(x for x in facts if x[8]=='采购公告')
    assert purchase[9:13]==('open','2099-12-31 18:00:00','buyer','公开采购')
    assert book['客户线索']['A1'].fill.fgColor.rgb.endswith('333333')


def test_feedback_invalidates_purchase_and_archived_lead_cannot_be_saved(client):
    """client为隔离客户端；负面反馈使公告待核状态可见，归档主体不被自动入库复活。"""
    identifier,candidate,rows,job=fixture(client)
    client.post('/api/leads/'+identifier+'/feedback',json={'kind':'stale'})
    row=client.get('/api/leads/'+identifier).json()
    assert row['qualification']['signal_score']==0 and row['procurement_signals'][0]['display_status']=='失效'
    identifier,candidate,rows,job=fixture(client)
    client.delete('/api/leads/'+identifier)
    with pytest.raises(ServiceError,match='归档'):save_candidate(candidate,rows,job)
