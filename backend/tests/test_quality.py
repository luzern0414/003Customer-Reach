"""方案 07：评分、截止时间、采购方补全、重复单位合并、筛选合批、完成判定与导出。"""
import asyncio
import io
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import select

from app import crawl, llm, services
from app.db import Company, Conversation, Fact, Job, JobCompany, Lead, Message, Notice, NoticeParty, Session
from app.research import quality, store
from app.research.pipeline import Run
from app.research.schemas import Contact, ListResult, ListRow, NoticeExtract, Party
from tests.test_core import signed_in

TODAY = date(2026, 10, 29)
TEXT = ('某市人民医院UPS电源采购项目招标公告 采购人信息 名称：某市人民医院 地址：某市健康路1号 联系方式：0791-1234567 '
        '项目所在地：江西省某市 投标截止时间：2026年11月13日 09点30分 预算金额：1200000元 采购需求：UPS电源2台')


def _job(title='质量测试') -> str:
    with Session.begin() as db:
        conversation = Conversation(title=title)
        db.add(conversation)
        db.flush()
        message = Message(conversation_id=conversation.id, role='user', content=title)
        db.add(message)
        db.flush()
        job = Job(conversation_id=conversation.id, message_id=message.id, plan={'plan': {
            'mode': 'find_buyers', 'title': title, 'product': 'UPS', 'keywords': ['UPS'], 'days': 90}})
        db.add(job)
        db.flush()
        return job.id


def _notice(url: str, buyer: str = '某市人民医院', text: str = TEXT, **update) -> dict:
    page = crawl.Page(url=url, status='captured', title='招标公告', text=text)
    extract = NoticeExtract(is_notice=True, title='某市人民医院UPS电源采购项目招标公告', kind='tender', published_on='2026-10-27',
                            deadline='2026年11月13日 09点30分', region='江西省某市', budget=1200000, budget_quote='预算金额：1200000元',
                            items='UPS电源2台', items_quote='采购需求：UPS电源2台', buyer=Party(name=buyer, quote=f'名称：{buyer}'),
                            buyer_contact=Contact(phone='0791-1234567', address='某市健康路1号',
                                                  quote='地址：某市健康路1号 联系方式：0791-1234567')).model_copy(update=update)
    return store.save_notice(extract, store.save_source(page, 'ccgp', 'notice', 'job'), page, {})


def test_parse_deadline_formats():
    assert quality.parse_deadline('2026年11月13日 09点30分') == '2026-11-13 09:30'
    assert quality.parse_deadline('2026-11-17 09:30') == '2026-11-17 09:30'
    assert quality.parse_deadline('2026年11月18日 09时30分00秒') == '2026-11-18 09:30'
    assert quality.parse_deadline('2026年11月3日上午9时00分（北京时间）') == '2026-11-03 09:00'
    assert quality.parse_deadline('2026年11月3日下午2点') == '2026-11-03 14:00'
    assert quality.parse_deadline('2026年9月15日') == '2026-09-15'
    assert quality.parse_deadline('详见招标文件') == '' and quality.parse_deadline('2026年13月40日') == ''
    assert quality.days_left('2026-11-03 09:00', TODAY) == 5 and quality.days_left('', TODAY) is None


def test_score_spreads_by_deadline_budget_contact():
    score = lambda **kw: quality.score_lead(**{'kind': 'tender', 'published_on': '2026-10-27', 'today': TODAY, **kw})
    rich = score(deadline_at='2026-11-20', budget=4_750_000, has_contact=True)
    plain = score(deadline_at='2026-11-20')
    soon = score(deadline_at='2026-10-30')
    passed = score(deadline_at='2026-10-20', budget=4_750_000, has_contact=True)
    assert rich == 100 and plain == 85 and soon == 75 and passed < plain
    assert quality.score_lead('tender', '2026-10-27', today=TODAY) > quality.score_lead('award', '2026-10-27', today=TODAY)
    assert quality.score_lead('tender', '2020-01-01', today=TODAY) == 10


def test_save_notice_fills_buyer_region_phone_address():
    saved = _notice('https://www.ccgp.gov.cn/q/fill.htm', buyer='补全测试医院', text=TEXT.replace('某市人民医院', '补全测试医院'))
    with Session() as db:
        company = db.get(Company, saved['buyer_id'])
        notice = db.get(Notice, saved['notice_id'])
        facts = {f.field: f for f in db.scalars(select(Fact).where(Fact.company_id == company.id))}
    assert notice.deadline_at == '2026-11-13 09:30'
    assert (company.region, company.phone, company.address) == ('江西省某市', '0791-1234567', '某市健康路1号')
    assert facts['phone'].quote.endswith('0791-1234567') and facts['phone'].source_id == notice.source_id


def test_backfill_and_merge_duplicates_are_idempotent():
    saved = _notice('https://www.ccgp.gov.cn/q/merge-a.htm', buyer='合并测试学院', text=TEXT.replace('某市人民医院', '合并测试学院'))
    with Session.begin() as db:
        # 旧版本留下的数据：截止时间未规范化、“本级”单位单独成一家且各有线索
        db.get(Notice, saved['notice_id']).deadline_at = ''
        dup = Company(name='合并测试学院本级')
        db.add(dup)
        db.flush()
        db.add(NoticeParty(notice_id=saved['notice_id'], company_id=dup.id, role='buyer', quote='名称：合并测试学院本级'))
        db.add(Lead(company_id=dup.id, notice_id=saved['notice_id'], score=99, status='contacted', owner='王五'))
        db.add(Lead(company_id=saved['buyer_id'], notice_id=saved['notice_id'], score=50))
        dup_id = dup.id
    stats = quality.backfill()
    assert stats['merged'] >= 1 and stats['deadlines'] >= 1
    with Session() as db:
        assert db.get(Company, dup_id) is None
        lead = db.scalar(select(Lead).where(Lead.company_id == saved['buyer_id']))
        assert (lead.status, lead.owner) == ('contacted', '王五')
        assert db.get(Notice, saved['notice_id']).deadline_at == '2026-11-13 09:30'
    again = quality.backfill()
    assert again['merged'] == 0 and again['deadlines'] == 0 and again['renamed'] == 0


def test_screen_batches_across_lists(monkeypatch):
    job_id = _job('合批测试')
    lists = {}
    for i in range(3):
        page = crawl.Page(url=f'https://ccgp.gov.cn/batch/{i}', status='captured', title='列表',
                          text=' '.join(f'UPS电源采购公告{i}-{n} 2026-10-2{n % 9} 某单位{n}' for n in range(30)))
        links = [{'text': f'UPS电源采购公告{i}-{n}', 'url': f'https://www.ccgp.gov.cn/cggg/batch/{i}-{n}.htm'} for n in range(30)]
        links.append({'text': '同一份UPS电源采购公告', 'url': f'https://www.ccgp.gov.cn/cggg/batch/same-{i}.htm'})
        lists[f'ccgp|UPS|{i}'] = {'site': 'ccgp', 'query': 'UPS', 'status': 'captured', 'links': links,
                                  'source_id': store.save_source(page, 'ccgp', 'list', job_id)}
    run = Run(job_id)
    run.state['lists'] = lists
    calls, active, peak = [], [0], [0]

    async def fake_structured(schema, prompt, system=''):
        calls.append(prompt)
        active[0] += 1
        peak[0] = max(peak[0], active[0])
        await asyncio.sleep(0.05)
        active[0] -= 1
        n = prompt.count('\n[')
        # 含 2-29 的整批输出超长，拆成两半后能判断完
        if 'UPS电源采购公告2-29' in prompt and n > 6:
            raise ValueError('模型输出超出长度上限')
        return ListResult(rows=[ListRow(link=i, relevant=True, kind='tender') for i in range(n)])

    monkeypatch.setattr(llm, 'structured', fake_structured)
    asyncio.run(run.screen())
    # 91 个不同标题合成 5 批（20、20、20、20、11），同时判断；列表摘录随链接送给模型；
    # 输出超长的一批拆成两半再判断（共 7 次调用），全部链接都有结果
    assert len(calls) == 7 and peak[0] > 1 and '列表摘录：UPS电源采购公告0-0' in calls[0]
    assert sum(c.count('同一份UPS电源采购公告') for c in calls) == 1
    assert len(run.state['screened']) == 3 and not any('未能筛选' in p for p in run.problems)
    assert len(run.state['candidates']) == 91


def test_final_status_tolerates_few_read_failures():
    job_id = _job('完成判定')
    run = Run(job_id)
    run.state['read'] = {f'u{i}': 'saved' for i in range(19)} | {'bad': 'failed'}
    run.problems = ['中国政府采购网公告未读取：读取失败']
    assert run.final_status() == 'partial'  # 还没有结果
    saved = _notice('https://www.ccgp.gov.cn/q/final.htm', buyer='判定测试医院', text=TEXT.replace('某市人民医院', '判定测试医院'))
    store.link_job(job_id, saved['buyer_id'], 'buyer')
    assert run.final_status() == 'completed'
    run.problems.append('已达到本次任务的时间上限')
    assert run.final_status() == 'partial'
    run.problems = [f'站点{i}公告未读取：读取失败' for i in range(6)]
    run.state['read'] = {f'u{i}': 'saved' for i in range(14)} | {f'bad{i}': 'failed' for i in range(6)}
    assert run.final_status() == 'partial'


def test_leads_open_filter_report_reuse_and_export():
    open_one = _notice('https://www.ccgp.gov.cn/q/open.htm', buyer='可投标医院', text=TEXT.replace('某市人民医院', '可投标医院'))
    closed = _notice('https://www.ccgp.gov.cn/q/closed.htm', buyer='已截止医院',
                     text=TEXT.replace('某市人民医院', '已截止医院').replace('2026年11月13日', '2020年1月2日'),
                     deadline='2020年1月2日 09点30分')
    for saved in (open_one, closed):
        store.upsert_lead(saved['buyer_id'], saved['notice_id'], 'c1', '招标公告：测试')
    client = signed_in()
    names = lambda params: {i['company']['name'] for i in client.get('/api/leads', params={**params, 'page_size': 100}).json()['items']}
    assert {'可投标医院', '已截止医院'} <= names({})
    assert '可投标医院' in names({'open': 1}) and '已截止医院' not in names({'open': 1})
    item = next(i for i in client.get('/api/leads', params={'q': '已截止医院'}).json()['items'])
    assert item['notice']['expired'] and item['notice']['days_left'] < 0 and item['company']['phone'] == '0791-1234567'
    # 研究报告计入复用的已读公告
    job_id = _job('复用统计')
    with Session.begin() as db:
        job = db.get(Job, job_id)
        job.plan = {**job.plan, 'read': {'https://www.ccgp.gov.cn/q/open.htm': 'reused'}}
        db.add(JobCompany(job_id=job_id, company_id=open_one['buyer_id'], role='buyer'))
    report = services.job_report(job_id)
    assert report['counts']['notices'] == 1 and report['top'][0]['name'] == '可投标医院'
    # 导出补列并防公式注入
    with Session.begin() as db:
        db.get(Notice, open_one['notice_id']).items = '=HYPERLINK("x")'
    book = load_workbook(io.BytesIO(client.get('/api/leads/export', params={'q': '可投标医院'}).content))
    rows = list(book.active.iter_rows(values_only=True))
    assert rows[0][:9] == ('客户', '地区', '状态', '评分', '采购标的', '预算（万元）', '公告日期', '投标截止', '剩余天数')
    assert rows[1][4] == '\'=HYPERLINK("x")' and rows[1][5] == 120 and rows[1][10] == '0791-1234567'


def test_events_carry_time_for_resumed_timers():
    from app import events
    job_id = _job('事件时间')
    events.emit(job_id, 'tool.started', {'id': 't1', 'tool': 'screen', 'label': '筛选'})

    async def first() -> str:
        async def connected() -> bool:
            return False
        stream = events.sse(job_id, 0, connected)
        item = await stream.__anext__()
        await stream.aclose()
        return item

    item = asyncio.run(first())
    data = __import__('json').loads(item.split('data: ', 1)[1])
    assert data['id'] == 't1' and data['ts'].endswith('+00:00')
