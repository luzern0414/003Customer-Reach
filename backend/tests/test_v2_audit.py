"""独立审查发现的业务回归：主体关联、冲突、反馈、真实预算预留及SSE分页。"""
import asyncio
import uuid
import pytest
from sqlalchemy import select
from langchain_core.messages import HumanMessage, AIMessage
from langchain_core.outputs import ChatGeneration, LLMResult
from app.db import Session, Job, JobEvent, Evidence, ToolCall, serialize
from app.schemas import Candidate, Fact
from app.service import checked_candidate, save_candidate, ServiceError
from app.knowledge import record_feedback
from app.model_gateway import ModelBudget
from app.research_tools import ToolBroker


def task(client):
    """client为隔离HTTP客户端；创建真实任务并领取，返回ID。"""
    conv=client.post('/api/conversations').json()['id']
    job=client.post(f'/api/conversations/{conv}/messages',json={'content':'最多1家深圳灯具店','client_token':uuid.uuid4().hex}).json()['id']
    with Session.begin() as db:db.get(Job,job).status='running'
    return job


def proposal(name, phone='0755-11111111', body=None, role='retailer'):
    """name为隔离主体，phone为公开fixture值，body可覆写；返回逐字有出处的候选/资料。"""
    body=body or f'{name}是深圳零售门店，电话{phone}，主营吊灯。'
    sources={'audit':{'id':'audit','body':body,'title':'隔离审查来源','url':'https://example.com/audit','kind':'目录','captured_at':'2026-10-01 00:00:00'}}
    candidate=Candidate(display_name=name,city='深圳',buyer_role=role,matched_products=['吊灯'],match_reason='测试',facts=[
        {'field':'name','value':name,'quote':body,'source_id':'audit'},
        {'field':'buyer_role','value':'零售门店' if role=='retailer' else '代理','quote':body,'source_id':'audit'},
        {'field':'phone','value':phone,'quote':body,'source_id':'audit'}])
    return candidate,sources


def test_same_page_is_not_same_entity_and_agency(client):
    candidate,sources=proposal('深圳甲制造公司',body='深圳甲制造公司制造灯具。深圳乙灯饰店是零售门店，电话0755-11111111。')
    candidate.facts[1].quote='深圳乙灯饰店是零售门店，电话0755-11111111。'
    with pytest.raises(ServiceError,match='主体'):checked_candidate(candidate,sources)
    candidate.facts[1].quote=sources['audit']['body']
    with pytest.raises(ServiceError,match='主体'):checked_candidate(candidate,sources)
    candidate,sources=proposal('深圳甲制造公司',body='深圳甲制造公司是制造商，深圳乙灯饰店是零售门店，电话0755-11111111。')
    with pytest.raises(ServiceError,match='主体'):checked_candidate(candidate,sources)
    candidate,sources=proposal('深圳采购代理机构',body='深圳采购代理机构负责采购代理，电话0755-11111111。',role='dealer')
    with pytest.raises(ServiceError,match='不是产品经销商'):checked_candidate(candidate,sources)


def test_conflicting_automatic_phone_preserved(client):
    candidate,sources=proposal('深圳冲突审查灯具店')
    lead=save_candidate(candidate,sources,'no-job')
    changed,other=proposal(candidate.display_name,'0755-22222222')
    assert save_candidate(changed,other,'no-job')['id']==lead['id']
    phone=client.get('/api/leads/'+lead['id']).json()['facts']['phone']
    assert phone['value']=='0755-11111111' and phone['verification']=='disputed'
    assert phone['alternatives'][0]['value']=='0755-22222222'


@pytest.mark.parametrize('kind',['incorrect','stale'])
def test_feedback_blocks_old_facts_even_after_cache(client,kind):
    candidate,sources=proposal('深圳'+kind+'审查灯具店')
    lead=save_candidate(candidate,sources,'no-job')
    record_feedback(lead['id'],kind,'隔离纠错')
    with pytest.raises(ServiceError,match='已被用户'):save_candidate(candidate,sources,'no-job')
    assert client.get('/api/leads/'+lead['id']).json()['score']==0


def test_historical_procurement_not_current_signal(client):
    candidate,sources=proposal('深圳历史审查灯具店',body='深圳历史审查灯具店是零售门店，电话0755-11111111。2025年采购吊灯已中标。')
    candidate.facts.append(Fact(field='procurement',value='采购吊灯',quote=sources['audit']['body'],source_id='audit'))
    lead=save_candidate(candidate,sources,'no-job')
    assert client.get('/api/leads/'+lead['id']).json()['qualification']['signal_score']==0


def test_model_task_budget_failure_usage_and_resume(client):
    job=task(client)
    async def run():
        first=ModelBudget(job,{'max_model_calls':2,'max_model_tokens':20000})
        await first.on_chat_model_start({},[[HumanMessage(content='中文灯具')]],run_id='failed')
        # 未收到usage/失败仍占额度；新实例模拟恢复，不能将累计账本清零。
        second=ModelBudget(job,{'max_model_calls':2,'max_model_tokens':20000})
        await second.on_chat_model_start({},[[HumanMessage(content='第二次')]],run_id='ok')
        await second.on_llm_end(LLMResult(generations=[[ChatGeneration(message=AIMessage(content='结果',usage_metadata={'input_tokens':10,'output_tokens':5,'total_tokens':15}))]]),run_id='ok')
        with pytest.raises(ServiceError,match='模型预算'):await second.on_chat_model_start({},[[HumanMessage(content='第三次')]],run_id='blocked')
    asyncio.run(run())
    metrics=client.get('/api/jobs/'+job).json()['metrics']['model_budget']
    assert metrics['calls']==2 and metrics['reported_tokens']==15 and metrics['estimated_calls']==1


def test_concurrent_reads_reserve_one_slot(client,monkeypatch):
    job=task(client);calls=[]
    async def fetch(self,args):
        """self为真实Broker，args为请求fixture；短等待制造并发预留条件。"""
        calls.append(args);await asyncio.sleep(.02);return {'sources':[]}
    monkeypatch.setattr(ToolBroker,'fetch',fetch)
    options={'max_searches':1,'max_reads':1,'enabled_channels':['official']}
    async def run():
        broker=ToolBroker(job,options,{})
        results=await asyncio.gather(*(broker.execute('fetch_source',{'url':f'https://example.com/{n}','channel':'official'}) for n in (1,2)))
        assert sum(x.get('error')=='budget_exhausted' for x in results)==1
    asyncio.run(run());assert len(calls)==1
    with Session() as db:assert len(list(db.scalars(select(ToolCall).where(ToolCall.job_id==job))))==1


def test_sse_drains_pages_before_terminal_and_replays_cursor(client):
    job=task(client)
    with Session.begin() as db:
        db.get(Job,job).status='completed'
        for index in range(320):db.add(JobEvent(job_id=job,event_type='message.delta',stage='answer',message='',payload={'delta':str(index)}))
        db.flush();ids=list(db.scalars(select(JobEvent.id).where(JobEvent.job_id==job).order_by(JobEvent.id)))
    stream=client.get('/api/jobs/'+job+'/events').text
    assert stream.count('event: message.delta')==320 and stream.count('event: status')==1
    assert stream.index('"delta": "319"')<stream.index('event: status')
    replay=client.get('/api/jobs/'+job+'/events?after='+str(ids[149]),headers={'Last-Event-ID':str(ids[199])}).text
    assert replay.count('event: message.delta')==120 and '"delta": "199"' not in replay


def test_saved_at_cap_resumes_pending_enrichment(client,monkeypatch):
    from app import agent as graph_module, research_tools
    from app.schemas import ProductProfile, Product, ResearchPlan, ResearchStep, CandidateBatch
    from app.research_tools import Paused
    from langchain_core.messages import AIMessageChunk
    import json
    job=task(client);name='深圳断点补全审查灯具店';paused=[False];extracts=[0];network=[]
    body=f'{name}是深圳零售门店，电话0755-11111111。地址深圳市南山区。主营吊灯。'
    async def provider(tool,args):
        """tool/args为搜索fixture，记录实际Broker调用，不发真实网络。"""
        network.append(args);return 'Title: 断点补全fixture\nURL: https://example.com/resume-enrich-'+str(len(network))+'\n'+body
    async def read(state):
        """state为真实图；达到客户数量上限后故意暂停待补全节点，恢复必须再次进入本节点。"""
        if state.get('saved_ids') and not paused[0]:paused[0]=True;raise Paused()
        return {}
    async def output(state,schema,prompt,label):
        """state/schema/prompt/label为真实结构化节点参数；只替换外部模型输出。"""
        if schema is ProductProfile:return ProductProfile(products=[Product(name='吊灯')],region='深圳',summary='隔离断点画像')
        if schema is ResearchPlan:return ResearchPlan(target_roles=['retailer'],steps=[
            ResearchStep(channel='official',query='深圳断点官网'+job,purpose='发现'),ResearchStep(channel='directory',query='深圳断点目录'+job,purpose='补全')])
        extracts[0]+=1
        candidate,_=proposal(name,body=body)
        with Session() as db:source_id=db.scalar(select(Evidence.id).where(Evidence.id.in_(state['source_ids']),Evidence.body.contains(name)).limit(1))
        for fact in candidate.facts:fact.source_id=source_id
        if extracts[0]>1:
            candidate.facts.extend([Fact(field='address',value='深圳市南山区',quote=body,source_id=source_id),Fact(field='categories',value='吊灯',quote=body,source_id=source_id)])
        return CandidateBatch(candidates=[candidate],missing_fields=[] if extracts[0]>1 else ['address','categories'])
    class Provider:
        """只替换真实provider协议，图路由和checkpoint仍用生产实现。"""
        def bind_tools(self,*args,**kwargs):
            """args/kwargs为绑定schema及选择参数；返回隔离provider自身。"""
            return self
        async def ainvoke(self,messages):
            """messages为明确搜索步骤，生成对应工具调用。"""
            step=json.loads(messages[-1].content.split('\n')[-1]);return AIMessage(content='',tool_calls=[{'id':uuid.uuid4().hex,'name':'SearchAccounts','args':{'query':step['query'],'channel':step['channel']}}])
        async def astream(self,messages):
            """messages为结果上下文；返回单块测试结果，不冒充真实模型流式验收。"""
            yield AIMessageChunk(content='已核验该测试客户。')
    monkeypatch.setattr(research_tools,'call_exa',provider);monkeypatch.setattr(graph_module,'read_node',read)
    monkeypatch.setattr(graph_module,'cached_model',output);monkeypatch.setattr(graph_module,'model',lambda:Provider())
    client.patch('/api/settings',json={'max_leads':1,'max_searches':3,'max_reads':0})
    asyncio.run(graph_module.run_job(job))
    before=client.get('/api/jobs/'+job).json();assert before['status']=='paused' and len(before['lead_ids'])==1
    assert client.get('/api/jobs/'+job+'/state').json()['next']==['read']
    with Session.begin() as db:db.get(Job,job).status='running'
    asyncio.run(graph_module.run_job(job))
    after=client.get('/api/jobs/'+job).json();assert after['status']=='completed' and after['lead_ids']==before['lead_ids']
    assert len(network)==2 and extracts[0]==2
    assert client.get('/api/leads/'+after['lead_ids'][0]).json()['facts']['address']['value']=='深圳市南山区'


def test_resume_does_not_reset_elapsed_budget(client):
    from app.agent import run_job
    job=task(client)
    client.patch('/api/settings',json={'max_seconds':60})
    with Session.begin() as db:db.get(Job,job).elapsed_ms=61000
    asyncio.run(run_job(job))
    row=client.get('/api/jobs/'+job).json()
    assert row['status']=='failed' and not row['tools']


def test_all_saved_customers_enriched_at_cap_and_disputed_gap(client):
    from app.agent import save_node,gap_node
    from app.schemas import SettingsInput
    from app.db import Lead
    job=task(client);saved=[];items=[];sources=[]
    for index in (1,2):
        name=f'深圳批量补全审查{index}灯具店'
        candidate,rows=proposal(name)
        with Session.begin() as db:
            source=Evidence(job_id=job,url=f'https://example.com/batch-{index}',title='隔离来源',body=rows['audit']['body']+'地址深圳市南山区。主营吊灯。',kind='测试',channel='official')
            db.add(source);db.flush();source_id=source.id;body=source.body
        rows[source_id]={**rows['audit'],'id':source_id,'body':body}
        for fact in candidate.facts:fact.source_id=source_id
        saved.append(save_candidate(candidate,rows,job)['id']);sources.append(source_id)
        candidate.facts.extend([Fact(field='address',value='深圳市南山区',quote=body,source_id=source_id),Fact(field='categories',value='吊灯',quote=body,source_id=source_id)])
        items.append(candidate.model_dump())
    options=SettingsInput(max_leads=2,max_searches=3).model_dump()
    state={'job_id':job,'options':options,'profile':{'products':[{'name':'吊灯'}],'region':'深圳'},'plan':{'target_roles':['retailer']},'saved_ids':saved,'candidates':items,'source_ids':sources,'gaps':[]}
    result=asyncio.run(save_node(state));assert result['saved_ids']==saved
    for lead in saved:assert client.get('/api/leads/'+lead).json()['facts']['address']['value']=='深圳市南山区'
    with Session.begin() as db:
        row=db.get(Lead,saved[1]);row.facts={**row.facts,'phone':{**row.facts['phone'],'verification':'disputed'}}
    steps=[{'channel':'official','query':'深圳原查询','purpose':'发现'},{'channel':'directory','query':'深圳下渠道','purpose':'补全'}]
    result=asyncio.run(gap_node({**state,'cursor':0,'steps':steps,'attempts':[]}))
    assert result['cursor']==1 and items[1]['display_name'] in result['steps'][1]['query']


def test_role_protocol_corrects_once_without_changing_fact(client,monkeypatch):
    from app import agent
    from app.schemas import SettingsInput,RoleCorrection
    job=task(client);name='深圳角色纠正审查照明公司';body=name+'为灯具销售点，电话0755-11111111。'
    with Session.begin() as db:
        source=Evidence(job_id=job,url='https://example.com/role-correct',title='销售点',body=body,kind='测试',channel='directory')
        db.add(source);db.flush();source_id=source.id
    candidate=Candidate(display_name=name,city='深圳',buyer_role='retailer',matched_products=['吊灯'],match_reason='协议fixture',facts=[
        Fact(field='name',value=name,quote=body,source_id=source_id),Fact(field='buyer_role',value='销售点',quote=body,source_id=source_id)])
    corrections=[]
    async def corrected(state,schema,prompt,label):
        """state/schema/prompt/label为纠正请求；仅返回角色枚举，不改事实或资料来源。"""
        assert schema is RoleCorrection;corrections.append(label);return RoleCorrection(buyer_role='dealer')
    monkeypatch.setattr(agent,'cached_model',corrected)
    result=asyncio.run(agent.save_node({'job_id':job,'cursor':0,'options':SettingsInput().model_dump(),'profile':{'products':[{'name':'吊灯'}]},'plan':{'target_roles':['dealer','retailer']},'candidates':[candidate.model_dump()],'source_ids':[source_id],'saved_ids':[],'gaps':[]}))
    assert len(corrections)==1 and len(result['saved_ids'])==1
    detail=client.get('/api/leads/'+result['saved_ids'][0]).json()
    assert detail['buyer_role']=='dealer' and detail['facts']['buyer_role']['value']=='销售点'


def test_role_filtered_export_matches_crm(client):
    from io import BytesIO
    from openpyxl import load_workbook
    listed=client.get('/api/leads?buyer_role=dealer&page_size=100').json()
    book=load_workbook(BytesIO(client.get('/api/leads/export?buyer_role=dealer').content))
    rows=list(book['客户线索'].iter_rows(min_row=2,values_only=True))
    headers=[cell.value for cell in book['客户线索'][1]]
    assert len(rows)==listed['total'] and all(row[headers.index('买家角色')]=='dealer' for row in rows)
