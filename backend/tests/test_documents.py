"""真实二进制文档与本地OCR/LibreOffice回归，不用伪造OCR返回证明扫描件支持。"""
import io
import asyncio
import subprocess
from datetime import datetime
from pathlib import Path
import pytest
from app.config import ROOT
from app.document_parser import parse_document, _executable
from app.service import ServiceError


def make_documents(directory: Path) -> dict:
    """directory为项目内测试目录，创建可复核的实际文件；返回文件路径，不生成客户事实。"""
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.shared import Inches
    from openpyxl import Workbook
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas
    import xlwt
    directory.mkdir(parents=True, exist_ok=True)
    image = Image.new('RGB', (1400, 1000), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 48)
    for index, value in enumerate(['产品：星河轨道灯', '型号：CR-T20', '功率：20W', '单价：58元', '地区：深圳', '仅用于文档解析验证']):
        draw.text((90, 80 + index * 120), value, fill='black', font=font)
    image.save(directory / 'scan.png')
    image.save(directory / 'scan.pdf', 'PDF', resolution=150)
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = '解析验证产品目录'
    doc.sections[0].footer.paragraphs[0].text = '资料来自测试文件'
    doc.add_paragraph('目录开始：星河轨道灯')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'CR-T20'
    table.cell(0, 1).text = '20W / 58元'
    table.cell(1, 0).merge(table.cell(1, 1)).text = '合并单元格说明'
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = '嵌套规格：4000K'
    doc.add_paragraph('目录结束：深圳经销商')
    doc.add_picture(str(directory / 'scan.png'), width=Inches(4))
    doc.save(directory / 'catalog.docx')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '产品'
    sheet.append(['名称', '型号', '功率', '单价'])
    sheet.append(['星河轨道灯', 'CR-T20', '20W', 58])
    sheet['A3'] = datetime(2026, 10, 2, 9, 30)
    sheet['B3'] = True
    sheet['C3'] = '=D2*2'
    sheet['D3'] = '#DIV/0!'
    sheet.merge_cells('A5:B5')
    sheet['A5'] = '合并说明'
    workbook.create_sheet('目标地区')['A1'] = '深圳'
    workbook.save(directory / 'catalog.xlsx')
    workbook.close()
    blank = Workbook()
    blank.save(directory / 'empty.xlsx')
    blank.close()
    book = xlwt.Workbook()
    old_sheet = book.add_sheet('产品')
    for col, value in enumerate(['星河轨道灯', 'CR-T20', 58, True, '20W']):
        old_sheet.write(0, col, value)
    old_sheet.write(1, 0, datetime(2026, 10, 2), xlwt.easyxf(num_format_str='YYYY-MM-DD'))
    book.save(str(directory / 'catalog.xls'))
    pdf = canvas.Canvas(str(directory / 'text.pdf'))
    pdf.drawString(70, 740, 'Track light CR-T20 / 20W / price 58 CNY')
    pdf.save()
    hybrid=canvas.Canvas(str(directory / 'single-page-mixed.pdf'))
    hybrid.drawString(70,790,'Product catalog heading')
    hybrid.drawImage(str(directory/'scan.png'),40,110,width=510,height=600)
    hybrid.save()
    mixed = PdfWriter()
    for file in ['text.pdf', 'scan.pdf']:
        for page in PdfReader(str(directory / file)).pages:
            mixed.add_page(page)
    mixed.write(str(directory / 'mixed.pdf'))
    encrypted = PdfWriter()
    encrypted.add_page(PdfReader(str(directory / 'text.pdf')).pages[0])
    encrypted.encrypt('fixture-only')
    encrypted.write(str(directory / 'encrypted.pdf'))
    profile = directory / 'fixture-lo'
    profile.mkdir(exist_ok=True)
    subprocess.run([_executable('soffice'), '-env:UserInstallation=' + profile.as_uri(), '--headless',
        '--convert-to', 'doc:MS Word 97', '--outdir', str(directory), str(directory / 'catalog.docx')],
        check=True, capture_output=True, timeout=90)
    assert (directory / 'catalog.doc').is_file(), 'LibreOffice必须实际生成DOC，不能只看退出码'
    return {x.name: x for x in directory.iterdir() if x.is_file()}


@pytest.fixture(scope='module')
def documents():
    return make_documents(ROOT / '.codex-tmp/document-fixtures')


def test_word_order_and_nested_tables(documents):
    result = parse_document('.docx', documents['catalog.docx'].read_bytes())
    assert result.text.index('目录开始') < result.text.index('CR-T20') < result.text.index('目录结束')
    assert '4000K' in result.text and '解析验证产品目录' in result.text and '资料来自测试文件' in result.text
    assert result.text.count('合并单元格说明') == 1
    assert result.metadata['tables'] == 2
    assert result.metadata['images_ocr'] == 1 and '插图OCR' in result.text


@pytest.mark.parametrize('nested',[False,True])
def test_empty_word_tables_do_not_become_document_content(nested,client):
    """nested控制真实DOCX空嵌套表，client为隔离API；人工表格标记不能使空文件上传成功。"""
    from docx import Document
    from io import BytesIO
    document=Document()
    table=document.add_table(rows=1,cols=1)
    if nested:table.cell(0,0).add_table(rows=1,cols=1)
    buffer=BytesIO();document.save(buffer)
    with pytest.raises(ServiceError,match='没有可分析'):parse_document('.docx',buffer.getvalue())
    response=client.post('/api/attachments',files={'file':('empty-table.docx',buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code==400 and '没有可分析' in response.json()['error']


def test_single_page_text_and_scanned_body(documents):
    result=parse_document('.pdf',documents['single-page-mixed.pdf'].read_bytes())
    assert 'Product catalog heading' in result.text and 'CR-T20' in result.text and '58' in result.text
    assert result.metadata['ocr_pages']==[1]
    assert result.metadata['page_details'][0]['method']=='text+ocr'


def test_legacy_word_conversion(documents):
    result = parse_document('.doc', documents['catalog.doc'].read_bytes())
    assert '星河轨道灯' in result.text and 'CR-T20' in result.text and '4000K' in result.text
    assert result.metadata['method'].startswith('LibreOffice')


def test_excel_coordinates_formula_and_dates(documents):
    result = parse_document('.xlsx', documents['catalog.xlsx'].read_bytes())
    assert 'B2: CR-T20' in result.text and 'D2: 58' in result.text
    assert 'C3: =D2*2 [无缓存值，未计算]' in result.text
    assert 'A3: 2026-10-02 09:30:00' in result.text and 'B3: TRUE' in result.text
    assert '#DIV/0!' in result.text and 'A5:B5' in result.text and '[工作表: 目标地区]' in result.text
    assert result.metadata['formula_without_cache'] == 1
    with pytest.raises(ServiceError, match='没有可分析'):
        parse_document('.xlsx', documents['empty.xlsx'].read_bytes())


def test_legacy_excel(documents):
    result = parse_document('.xls', documents['catalog.xls'].read_bytes())
    assert 'B1: CR-T20' in result.text and 'C1: 58' in result.text and 'D1: TRUE' in result.text
    assert '2026-10-02' in result.text and result.metadata['warnings']


@pytest.mark.parametrize('name,pages', [('text.pdf', []), ('scan.pdf', [1]), ('mixed.pdf', [2])])
def test_actual_pdf_and_chinese_ocr(documents, name, pages):
    result = parse_document('.pdf', documents[name].read_bytes())
    assert '20W' in result.text and '58' in result.text
    if pages:
        assert '轨道灯' in result.text and '深圳' in result.text
    assert result.metadata['ocr_pages'] == pages
    if name == 'mixed.pdf':
        assert '[第1页 · text]' in result.text and '[第2页 · ocr]' in result.text


def test_encrypted_and_truncated(documents):
    with pytest.raises(ServiceError, match='加密'):
        parse_document('.pdf', documents['encrypted.pdf'].read_bytes())
    result = parse_document('.txt', ('长文档' * 50000).encode())
    assert result.truncated and len(result.text) == 120000 and result.metadata['warnings']


@pytest.mark.parametrize('suffix', ['.pdf', '.docx', '.xlsx', '.xls'])
def test_corrupt_documents(suffix):
    with pytest.raises(ServiceError, match='无法解析'):
        parse_document(suffix, b'broken-file')


@pytest.mark.parametrize('name', ['catalog.docx', 'catalog.doc', 'catalog.xlsx', 'catalog.xls', 'text.pdf', 'scan.pdf', 'mixed.pdf'])
def test_upload_preview_and_persistence(client, documents, name):
    response = client.post('/api/attachments', files={'file': (name, documents[name].read_bytes())})
    assert response.status_code == 201, response.text
    meta = response.json()
    assert meta['parse']['elapsed_ms'] >= 0 and meta['characters'] > 0 and 'stored_path' not in meta
    detail = client.get('/api/attachments/' + meta['id']).json()
    assert '20W' in detail['text'] and 'stored_path' not in detail
    assert detail['parse'] == meta['parse']
    conv = client.post('/api/conversations').json()['id']
    job = client.post(f'/api/conversations/{conv}/messages', json={'content':'文档解析回归测试',
        'attachment_ids':[meta['id']], 'client_token':'document-'+meta['id']}).json()
    value = client.get('/api/conversations/'+conv).json()
    assert value['attachments'][0]['filename'] == name
    assert value['jobs'][0]['id'] == job['id']
    client.post('/api/jobs/'+job['id']+'/pause')


def test_preview_missing(client):
    assert client.get('/api/attachments/missing-document-id').status_code == 404


def test_historical_tools_and_elapsed(client):
    from app.db import Session, Job, ToolCall
    conv = client.post('/api/conversations').json()['id']
    from app import agent
    rows=[]
    for index in range(2):
        job=client.post(f'/api/conversations/{conv}/messages',json={'content':'你好','client_token':'history-'+str(index)+'-'+conv}).json()
        asyncio.run(agent.run_job(job['id']));rows.append(job)
    with Session.begin() as db:
        for index, job in enumerate(rows):
            db.get(Job, job['id']).elapsed_ms = 1200 + index * 1000
            db.add(ToolCall(job_id=job['id'], operation_id='test-op-'+job['id'],tool='search_accounts',
                channel='official', arguments={'query':'测试客户','api_key':'DO-NOT-EXPOSE'}, status='completed',
                result={'sources':[{'url':'https://example.com','title':'隔离结果','body':'DO-NOT-EXPOSE'}]},elapsed_ms=900))
    result = client.get('/api/conversations/'+conv).json()
    assert len(result['jobs']) == 2
    assert result['jobs'][0]['id'] == rows[-1]['id']
    for job in result['jobs']:
        assert job['elapsed_ms'] in {1200,2200} and job['reply_message_id']
        assert job['tools'][0]['source_count'] == 1
    assert 'DO-NOT-EXPOSE' not in str(result)
