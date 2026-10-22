"""公开微信独立保存/复制/导出，不从掩码或微信数字生成电话，不授予平台号码给客户。"""
import uuid
from io import BytesIO
import pytest
from openpyxl import load_workbook
from app.contact_view import wechat_valid,contact_view
from app.contact_owner import public_wechat


@pytest.mark.parametrize('value,valid',[
    ('15013617190',True),('test_wechat_01',True),('wxid_123456789012345',True),
    ('150136*****',False),('unknown',False),('微信同号',False),('测试昵称',False),
    ('1501361719',False),('',False),('15013617190 extra',False)])
def test_wechat_value_format_is_separate_from_phone(value,valid):
    """value为明文/掩码/未知测试值、valid为格式预期；有效微信不派生tel入口，电话缺口保持。"""
    assert wechat_valid(value)==valid
    view=contact_view({'wechat':{'value':value}},[])
    assert ('wechat' in view['usable_fields'])==valid
    assert not view['phone_href'] and 'phone' in view['missing_fields']


def test_wechat_requires_label_and_the_matching_owner_occurrence():
    """无入参；同值电话不能替平台微信授资格，重复联系块只接受真实同公司位置，不猜同号。"""
    name='深圳隔离微信商行';value='15013617190'
    body=name+'\n联系电话：'+value+'\n平台微信：'+value
    assert not public_wechat(name,value,'微信：'+value,body,name)
    assert not public_wechat(name,value,'联系电话：'+value,body,name)
    assert not public_wechat(name,value,'电话：'+value+'，微信同号',body,name)
    own=name+'\n其它联系方式：微信：'+value
    assert public_wechat(name,value,'其它联系方式：微信：'+value,own,name)
    duplicate=body+'\n'+own
    assert public_wechat(name,value,'微信：'+value,duplicate,name)
    foreign='其他隔离企业有限公司\n微信：'+value
    assert not public_wechat(name,value,'微信：'+value,foreign,name)


def setup(client):
    """client为隔离API；返回真实数据库候选/来源/任务ID，只含明确测试企业和明文微信，不请求公网。"""
    from app.db import Session,Job
    from app.gateway import archive_sources
    from app.schemas import Candidate,Fact
    conv=client.post('/api/conversations').json()['id']
    job=client.post('/api/conversations/'+conv+'/messages',json={'content':'找深圳灯具批发商','client_token':uuid.uuid4().hex}).json()['id']
    with Session.begin() as db:db.get(Job,job).status='running'
    name='深圳'+uuid.uuid4().hex[:6]+'隔离微信照明有限公司';value='test_wechat_01'
    role=name+'是一家照明灯具批发企业。'
    quote='业务联系 微信：'+value
    body=role+'\n地址：深圳隔离测试地址\n'+quote
    source=archive_sources(job,[{'url':'https://example.org/'+uuid.uuid4().hex,'title':name,'body':body}],'网页正文','official')[0]
    candidate=Candidate(display_name=name,city='深圳',buyer_role='dealer',match_reason='隔离微信业务验证',facts=[
        Fact(field='name',value=name,quote=name,source_id=source['id']),
        Fact(field='buyer_role',value='批发',quote=role,source_id=source['id']),
        Fact(field='wechat',value=value,quote=quote,source_id=source['id'])])
    return candidate,source,job


def test_wechat_full_save_projection_export_feedback_and_manual_clear(client):
    """client为隔离API；保存/排序/导出/反馈/人工覆盖全链，无电话值或动作，失效旧引文不可复活。"""
    from app.service import save_candidate,ServiceError
    from app.answer_summary import saved_summary,answer_context
    candidate,source,job=setup(client)
    key=save_candidate(candidate,{source['id']:source},job)['id']
    detail=client.get('/api/leads/'+key).json()
    assert detail['contact_summary']['usable_fields']==['wechat'] and detail['facts']['wechat']['copyable']
    assert not detail['facts'].get('phone') and detail['contact_summary']['purpose']=='公开业务联系'
    assert not detail['contact_summary']['phone_href']
    assert '微信：test_wechat_01' in saved_summary([detail])
    assert any('wechat' in card['fields'] for card in answer_context([detail])[0]['contact_evidence'])
    book=load_workbook(BytesIO(client.get('/api/leads/export?q='+candidate.display_name).content))
    headers=[cell.value for cell in book['客户线索'][1]]
    assert book['客户线索'].cell(2,headers.index('公开业务微信')+1).value=='test_wechat_01'
    assert book['客户线索'].cell(2,headers.index('公开业务电话')+1).value is None
    assert any(row[2]=='wechat' and row[3]=='test_wechat_01' for row in book['事实与口径'].iter_rows(min_row=2,values_only=True))
    assert any(row[2]=='wechat' and row[3]=='业务联系 微信：test_wechat_01' for row in book['逐字段证据'].iter_rows(min_row=2,values_only=True))
    client.post('/api/leads/'+key+'/feedback',json={'kind':'stale'})
    invalid=client.get('/api/leads/'+key).json()
    assert not invalid['facts']['wechat']['copyable'] and 'wechat' not in invalid['contact_summary']['usable_fields']
    with pytest.raises(ServiceError):save_candidate(candidate,{source['id']:source},job)
    changed=client.patch('/api/leads/'+key,json={'wechat':'new_test_wechat'}).json()
    assert changed['facts']['wechat']['verification'].startswith('人工') and 'wechat' in changed['manual_fields']
    client.patch('/api/leads/'+key,json={'wechat':''})
    # PATCH沿现有原行响应契约，联系投影由GET重新取得，不能用测试要求新增响应字段。
    cleared=client.get('/api/leads/'+key).json()
    assert not cleared['facts']['wechat']['value'] and not cleared['contact_summary']['usable_fields']


def test_wechat_conflicts_remove_action_and_other_company_cannot_supply_it(client):
    """client隔离API；异值保持disputed，不把其他公司的微信或仅电话字段当本公司微信。"""
    from app.service import save_candidate,checked_candidate,ServiceError
    from app.gateway import archive_sources
    from app.schemas import Fact
    candidate,source,job=setup(client)
    key=save_candidate(candidate,{source['id']:source},job)['id']
    name=candidate.display_name;quote='微信：other_test_wechat'
    other=archive_sources(job,[{'url':'https://example.net/'+uuid.uuid4().hex,'title':name,'body':name+'\n'+quote}],'网页正文','official')[0]
    changed=candidate.model_copy(update={'facts':[candidate.facts[0],candidate.facts[1],Fact(field='wechat',value='other_test_wechat',quote=quote,source_id=other['id'])]})
    save_candidate(changed,{source['id']:source,other['id']:other},job)
    detail=client.get('/api/leads/'+key).json()
    assert detail['facts']['wechat']['verification']=='disputed' and not detail['facts']['wechat']['copyable']
    assert 'wechat' not in detail['contact_summary']['usable_fields']
    foreign={**other,'body':'其他隔离主体有限公司\n'+quote}
    with pytest.raises(ServiceError,match='微信必须'):checked_candidate(changed,{source['id']:source,other['id']:foreign})


@pytest.mark.parametrize('field',['phone','wechat'])
def test_number_only_disclosed_as_wechat_never_becomes_phone(client,field):
    """client隔离API、field测试误映射和正确映射；同一数字仅微信标签公开时不能作电话号码。"""
    from app.service import checked_candidate,ServiceError
    from app.schemas import Fact
    candidate,source,job=setup(client)
    value='15013617190';quote='微信：'+value
    source={**source,'body':source['body']+'\n'+quote}
    altered=candidate.model_copy(update={'facts':[*candidate.facts[:2],Fact(field=field,value=value,quote=value if field=='phone' else quote,source_id=source['id'])]})
    if field=='phone':
        with pytest.raises(ServiceError,match='微信号当作'):checked_candidate(altered,{source['id']:source})
    else:
        facts,_=checked_candidate(altered,{source['id']:source});assert facts['wechat']['value']==value and 'phone' not in facts
