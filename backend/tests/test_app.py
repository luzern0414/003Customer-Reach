"""初代关键契约回归，独立临时数据库，绝不清空真实客户数据。"""
import os
import uuid
from pathlib import Path
import pytest

from fastapi.testclient import TestClient
from app.main import app
from app.db import Session, Job
from app.service import checked_candidate, ServiceError, save_candidate
from app.schemas import Candidate
from app.gateway import split_sources

def test_crm_persistence(client):
    row = client.post('/api/leads', json={'display_name':'回归用例客户','city':'深圳','phone':'0755-12345678'}).json()
    assert row['facts']['phone']['verification'].startswith('人工')
    assert client.patch('/api/leads/'+row['id'], json={'owner':'本地测试','status':'contacted'}).status_code == 200
    assert client.post('/api/leads/'+row['id']+'/followups', json={'content':'本地自动验证记录','kind':'note'}).status_code == 201
    detail = client.get('/api/leads/'+row['id']).json()
    assert detail['owner'] == '本地测试' and len(detail['followups']) == 1
    assert client.get('/api/leads?status=contacted&q=回归').json()['total'] == 1
    assert client.get('/api/leads/export').content.startswith(b'PK')
    assert client.delete('/api/leads/'+row['id']).json()['archived']
    assert client.get('/api/leads/'+row['id']).status_code == 404

def test_job_idempotency_and_pause(client):
    conv = client.post('/api/conversations').json()['id']
    payload = {'content':'测试吊灯客户画像','client_token':'isolated-token-0001'}
    first = client.post(f'/api/conversations/{conv}/messages', json=payload).json()
    assert client.post(f'/api/conversations/{conv}/messages', json=payload).json()['id'] == first['id']
    assert client.post(f'/api/conversations/{conv}/messages', json={**payload,'client_token':'second-token-0002'}).status_code == 409
    assert client.post('/api/jobs/'+first['id']+'/pause').json()['status'] == 'paused'
    assert 'event: status' in client.get('/api/jobs/'+first['id']+'/events').text
    assert client.post('/api/jobs/'+first['id']+'/resume').json()['status'] == 'queued'
    assert client.post('/api/jobs/'+first['id']+'/resume').status_code == 409
    with Session.begin() as db:
        db.get(Job, first['id']).status = 'cancelling'
    assert client.post(f'/api/conversations/{conv}/messages', json={**payload,'client_token':'third-token-0003'}).status_code == 409

def test_upload_and_settings(client):
    text = '产品：轨道灯\n规格：20W\n价格：未提供'.encode('utf-8')
    assert client.post('/api/attachments', files={'file':('产品.txt',text,'text/plain')}).json()['characters'] > 0
    assert client.post('/api/attachments', files={'file':('坏文件.pdf',b'broken','application/pdf')}).status_code == 400
    assert client.post('/api/attachments', files={'file':('程序.exe',b'not-executable')}).status_code == 415
    assert client.post('/api/attachments', files={'file':('空.txt',b'')}).status_code == 413
    result = client.patch('/api/settings',json={'company_name':'测试公司','max_leads':2}).json()
    assert result['company_name'] == '测试公司' and 'api_key' not in result
    assert client.patch('/api/settings',json={'api_key':'forbidden'}).status_code == 422
    assert client.patch('/api/settings',json={'max_searches':100}).status_code == 422
    assert client.post('/api/leads',json={'display_name':'  '}).status_code == 400
    assert client.post('/api/leads',json={'display_name':'客户','next_followup':'bad-time'}).status_code == 400

SOURCE = {'s1':{'id':'s1','body':'深圳测试灯具店。地址深圳市南山区。联系电话0755-12345678。主营吊灯。注册资本100万。零售价200元。经营品牌甲牌。',
 'url':'https://example.com/public','title':'隔离测试来源','captured_at':'2026-10-01 00:00:00','kind':'测试'}}

def candidate(extra=None):
    """extra为可选事实字典，无值只返回名称事实；仅用于隔离验证。"""
    facts = [{'field':'name','value':'深圳测试灯具店','quote':'深圳测试灯具店','source_id':'s1'}]
    if extra: facts.append({'source_id':'s1',**extra})
    return Candidate(display_name='深圳测试灯具店',city='深圳',facts=facts,matched_products=['吊灯'],match_reason='测试匹配')

@pytest.mark.parametrize('extra',[
 {'field':'phone','value':'13800000000','quote':'联系电话13800000000'},
 {'field':'revenue','value':'100万','quote':'注册资本100万','scope':'企业','effective_date':'2026'},
 {'field':'revenue','value':'100万','quote':'注册资本100万'},
 {'field':'purchase_price','value':'200元','quote':'零售价200元'},
 {'field':'supplier','value':'甲牌','quote':'经营品牌甲牌'},
 {'field':'industry','value':'深圳','quote':'地址深圳市南山区'},
 {'field':'industry','value':'吊灯','quote':'主营吊灯'},
])
def test_reject_fabrication(extra):
    with pytest.raises(ServiceError): checked_candidate(candidate(extra), SOURCE)


def test_industry_requires_semantic_classification():
    """原文行业分类可保存；即使地区被贴上行业标签，也不能充当行业。"""
    source={'s1':{**SOURCE['s1'],'body':SOURCE['s1']['body']+'所属行业：零售业。行业：深圳。'}}
    clean,_=checked_candidate(candidate({'field':'industry','value':'零售业','quote':'所属行业：零售业'}),source)
    assert clean['industry']['value']=='零售业'
    with pytest.raises(ServiceError,match='行业需要明确行业分类'):
        checked_candidate(candidate({'field':'industry','value':'深圳','quote':'行业：深圳'}),source)

def test_preserve_manual_and_dedupe(client):
    c = candidate({'field':'phone','value':'0755-12345678','quote':'联系电话0755-12345678'})
    row = save_candidate(c,SOURCE,'no-job')
    client.patch('/api/leads/'+row['id'],json={'phone':'0755-88888888','status':'contacted'})
    assert save_candidate(c,SOURCE,'no-job')['id'] == row['id']
    saved = client.get('/api/leads/'+row['id']).json()
    assert saved['facts']['phone']['value'] == '0755-88888888' and saved['status'] == 'contacted'
    assert saved['industry'] == ''

def test_source_formats():
    assert len(split_sources('Title: 官网\nURL: https://example.com\n正文')) == 1
    assert len(split_sources('# 官网\nURL: https://example.com\n正文')) == 1
    assert len(split_sources('{"results":[{"url":"https://example.com","text":"正文"}]}')) == 1
    assert split_sources('没有来源URL的文本') == []

def test_sample_import_and_rich_export(client):
    """在隔离库验证历史标记、幂等和三张证据工作表，避免污染真实CRM。"""
    from io import BytesIO
    from openpyxl import load_workbook
    first = client.post('/api/leads/import-sample').json()
    assert first['imported'] == 20
    assert client.post('/api/leads/import-sample').json()['imported'] == 0
    rows = client.get('/api/leads?page_size=100').json()['items']
    assert sum('历史研究样本' in row['tags'] for row in rows) == 20
    manual = client.post('/api/leads',json={'display_name':'=公式验证客户'}).json()
    book = load_workbook(BytesIO(client.get('/api/leads/export').content))
    assert book.sheetnames == ['客户线索','事实与口径','逐字段证据']
    assert any(row[1] == "'=公式验证客户" for row in book['客户线索'].iter_rows(min_row=2, values_only=True))
    assert book['逐字段证据'].max_row > 1
    client.delete('/api/leads/'+manual['id'])
