"""企业模块使用真实归档和隔离SQLite核验口径、幂等与API联动，不调用外部模型。"""
import io
import uuid
import pytest
from sqlalchemy import create_engine, select, func
from sqlalchemy.orm import sessionmaker


@pytest.fixture(scope='module')
def workspace():
    """无入参；建立专用测试数据库并替换当前模块会话，结束恢复原会话，绝不接正式CRM。"""
    from app import db, main, service, company_import, company_service
    # 先载入会持有Session别名的运行模块，再成组切换/恢复；否则临时库期间首次导入会污染后续图测试。
    from app import agent, intent, research_tools, source_service, research_dataset, company_validation, company_research, model_gateway, knowledge
    from app import relation_service, company_identity, counterparty_research
    from app.config import ROOT
    from fastapi.testclient import TestClient
    path = ROOT / '.codex-tmp/tests' / ('companies-' + uuid.uuid4().hex)
    path.mkdir(parents=True)
    engine = create_engine('sqlite:///' + str(path / 'test.db'), connect_args={'check_same_thread': False})
    session = sessionmaker(engine, expire_on_commit=False)
    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(db, 'engine', engine)
        for module in (db,main,service,company_import,company_service,agent,intent,research_tools,source_service,
                       research_dataset,company_validation,company_research,model_gateway,knowledge,relation_service,company_identity,counterparty_research):patch.setattr(module,'Session',session)
        patch.setattr(company_import, 'DATA', path / 'data')
        with TestClient(main.app) as client:
            response = client.post('/api/companies/import-history')
            assert response.status_code == 200, response.text
            yield client, session, response.json()
    engine.dispose()


def test_history_is_idempotent_without_fake_jobs(workspace):
    """验证重复导入不增事实/公告/来源且没有工具任务，并阻止跨资料裸来源ID碰撞。"""
    from app.db import Evidence, Job, ToolCall, Conversation
    from app.company_models import CompanyFact, ProcurementProject, ProcurementNotice, CompanyRelation, ResearchCompany
    client, session, imported = workspace
    models = (Evidence, Job, ToolCall, Conversation, CompanyFact, ProcurementProject, ProcurementNotice, CompanyRelation, ResearchCompany)
    with session() as db: before = [db.scalar(select(func.count()).select_from(model)) for model in models]
    assert client.post('/api/companies/import-history').status_code == 200
    with session() as db:
        assert before == [db.scalar(select(func.count()).select_from(model)) for model in models]
        assert db.scalar(select(func.count()).select_from(Job)) == 0
        assert db.scalar(select(func.count()).select_from(ToolCall)) == 0
        f001 = [row for row in db.scalars(select(Evidence)) if row.source_meta.get('source_id') == 'F001']
        assert len(f001) == 2 and len({row.url for row in f001}) == 2
    assert client.get('/api/companies').json()['total'] == 4
    for conv in imported['conversation_ids']:
        assert client.get('/api/companies', params={'conversation_id': conv}).json()['total'] in {1, 3}


def test_project_roles_dates_currency_and_relation_totals(workspace):
    """验证真实买方项目归并、未知币种、海能直接公告缺口与上下游交易方向，合计不成为对手方。"""
    client, _, imported = workspace
    details = [client.get('/api/companies/' + id).json() for id in imported['company_ids']]
    hanon, byd = details
    assert byd['procurement_summary'] == {'projects': 11, 'notices': 15, 'plans': 0, 'quantity_disclosed': 8}
    assert all(p['payload']['currency_confirmed'] is False for p in byd['projects'])
    assert any('候选' in p['stage'] for p in byd['projects'])
    assert hanon['procurement_summary']['projects'] == 0 and hanon['procurement_summary']['plans'] == 6
    assert any('本轮未找到' in gap for gap in hanon['gaps'])
    assert hanon['relation_summary']['supplier'] == {'count': 5, 'amount': 15320813.62, 'share': 17.24, 'period':'2025', 'unit':'元', 'currency':'CNY'}
    assert hanon['relation_summary']['customer'] == {'count': 5, 'amount': 44696738.97, 'share': 12.36, 'period':'2025', 'unit':'元', 'currency':'CNY'}
    assert all(r['counterparty_name'] != '合计' for r in hanon['relations'])
    services = [r for r in hanon['relations'] if r['relation_type'] == 'service_supplier']
    assert len(services) == 1 and services[0]['amount'] == 43813.2
    applications = [p for p in hanon['projects'] if p['record_kind'] == 'brand_application']
    assert applications and all(p['payload']['direct_trade'] is False for p in applications)
    assert hanon['header']['facts']['headcount']['numeric_value'] == 566
    assert hanon['header']['facts']['headcount']['as_of'] == '2026-06-30'
    assert '控股子公司' in hanon['header']['facts']['headcount']['note']
    assert isinstance(hanon['header']['facts']['industry']['raw_value'], str)
    assert '仪器' in hanon['header']['facts']['industry']['raw_value']
    assert '研发' in hanon['header']['facts']['business']['raw_value']
    assert hanon['header']['facts']['revenue']['numeric_value'] == 361850165.54
    assert '12-31' in hanon['header']['facts']['revenue']['report_period']


def test_api_uses_updated_database_and_export(workspace):
    """修改隔离库事实后API与XLSX同步变化，证明前端消费契约没有图片样例值；恢复后保留其他用例原资料。"""
    from app.company_models import CompanyFact
    from openpyxl import load_workbook
    client, session, imported = workspace
    id = imported['company_ids'][0]
    original = client.get('/api/companies/' + id).json()['header']['facts']['headcount']
    with session.begin() as db:
        row = db.get(CompanyFact, original['id']); row.raw_value = 777; row.numeric_value = 777
    try:
        assert client.get('/api/companies/' + id).json()['header']['facts']['headcount']['raw_value'] == 777
        response = client.get('/api/companies/' + id + '/export')
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        assert any(str(row[3]) == '777' for row in workbook['企业事实'].iter_rows(min_row=2, values_only=True))
    finally:
        with session.begin() as db:
            row = db.get(CompanyFact, original['id']); row.raw_value = original['raw_value']; row.numeric_value = original['numeric_value']


def test_company_research_submission_does_not_require_seller_products(workspace):
    """无产品资料即可排队指定公司研究，重发幂等且跨公司token拒绝；不存在主体404。"""
    client, _, imported = workspace
    body = {'client_token': uuid.uuid4().hex, 'content': '核验最新公开采购信息'}
    first = client.post('/api/companies/' + imported['company_ids'][0] + '/research', json=body)
    assert first.status_code == 202
    assert first.json()['query_mode'] == 'company_query' and first.json()['status'] == 'queued'
    second = client.post('/api/companies/' + imported['company_ids'][0] + '/research', json=body)
    assert second.status_code == 202 and second.json()['id'] == first.json()['id']
    assert client.post('/api/companies/' + imported['company_ids'][1] + '/research', json=body).status_code == 409
    assert client.get('/api/companies/missing').status_code == 404


def test_header_relation_counts_use_latest_distinct_counterparties(workspace):
    """同一供应商跨年与转载记录不能增加摘要家数，完成后仅移除本用例隔离新增行。"""
    from app.company_models import CompanyRelation
    client,session,imported=workspace;cid=imported['company_ids'][0];ids=[]
    with session.begin() as db:
        for year in ['2024','2025','2025']:
            row=CompanyRelation(unique_key=uuid.uuid4().hex,company_id=cid,counterparty_name='CTC Analytics AG',relation_type='supplier',report_period=year,unit='元',currency='CNY',amount=4289389.84,share=4.83)
            db.add(row);db.flush();ids.append(row.id)
    try:
        detail=client.get('/api/companies/'+cid).json()
        assert detail['header']['supplier_count']==5
        assert detail['relation_summary']['supplier']['count']==5
    finally:
        with session.begin() as db:
            for key in ids:db.delete(db.get(CompanyRelation,key))
