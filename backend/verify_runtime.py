"""本地运行验收：实际API/证据/Excel/配置隔离，不调用模型或发送销售消息。"""
import io
import re
import httpx
from openpyxl import load_workbook
from app.config import ROOT, credentials

def main():
    """核对运行中的本机服务，失败断言；只输出状态和数量，不输出密钥/公开联系人。"""
    with httpx.Client(base_url='http://127.0.0.1:8000',timeout=30) as client:
        health = client.get('/api/health').json()
        assert health['worker_online'] and health['model'] == 'gpt-6-luna'
        settings = client.get('/api/settings')
        secret = credentials()[1]
        assert secret not in settings.text and 'api_key' not in settings.json()
        rows = client.get('/api/leads?page_size=100').json()['items']
        citations = 0
        for row in rows:
            assert '非真实客户' not in row['display_name']
            for e in row['evidence']:
                if '历史研究样本' in row['tags']: continue
                evidence = client.get('/api/evidence/'+e['source_id'])
                evidence.raise_for_status()
                normalize = lambda x: re.sub(r'\s+','',x)
                assert normalize(e['quote']) in normalize(evidence.json()['body'])
                assert normalize(e['value']) in normalize(e['quote'])
                citations += 1
        response = client.get('/api/leads/export')
        response.raise_for_status()
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ['客户线索','事实与口径','逐字段证据']
        assert book['客户线索'].max_row == len(rows)+1
        for path in (ROOT/'log/web').glob('*.jsonl*'):
            assert secret not in path.read_text(encoding='utf-8')
        print({'health':'ok','model':'gpt-6-luna','leads':len(rows),'citations_checked':citations,
               'export_sheets':len(book.sheetnames),'secret_exposure':False})

if __name__ == '__main__':
    main()
