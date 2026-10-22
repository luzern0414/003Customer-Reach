"""企业工作台投影与持久服务；业务值均由数据库事实和来源生成。"""
import hashlib
import io
import json
import re
from copy import copy
from decimal import Decimal
from sqlalchemy import select, func, or_, case
from .db import Session, Evidence, Conversation, Job, Message, Lead, now, serialize
from .company_models import Company, CompanyFact, ProcurementProject, ProcurementNotice, CompanyRelation, ResearchCompany
from .service import ServiceError, submit
from .contact_view import phone_info, website_link
from .source_service import entry_references
from .tender_service import tender_view, tender_source_ids


def stable_key(*parts) -> str:
    """parts为可序列化身份片段，可含空值；返回确定性SHA-256，不作主体相似合并。"""
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def canonical_field(label: str) -> str:
    """label为原披露字段名，空值返回detail；只映射语义明确字段，不推断行业或联系人用途。"""
    if label in {'营业收入', '营业总收入'}: return 'revenue'
    if label in {'员工总计期末人数', '在职员工总人数', '员工人数', '员工总数'}: return 'headcount'
    if '统一社会信用代码' in label: return 'credit_code'
    if re.search(r'^(?:所处)?行业(?:分类|归属)?(?:[（(].*)?$', label): return 'industry'
    if re.search(r'产品系列|主要产品|主营产品', label) and '服务项目' not in label: return 'products'
    if re.search(r'主要产品与服务项目|^主营业务(?!收入|成本|毛利|利润)|^经营范围|^业务概要|^业务性质$', label): return 'business'
    if re.search('公司网址|官方网站|官网网址|官网$', label): return 'website'
    if re.search('邮箱|电子信箱|电子邮件', label): return 'email'
    if re.search('电话|热线|传真', label): return 'phone' if '传真' not in label else 'fax'
    if '地址' in label: return 'address'
    if re.search('证券代码|股票代码', label): return 'stock_code'
    if re.search('证券简称|股票简称|公司简称', label): return 'short_name'
    if re.search('采购模式|采购组织|供应商管理|供应商准入|付款|采购品类|采购机制', label): return 'procurement_rule'
    return 'detail'


def ensure_company(db, name: str, **values) -> Company:
    """db为当前事务，name为证据原名且不能为空，values为已核验元数据；精确原名幂等，不模糊合并。"""
    name = name.strip()
    if not name: raise ServiceError('企业名称不能为空')
    key = stable_key('name', name)
    row = db.scalar(select(Company).where(Company.identity_key == key))
    if not row:
        row = Company(identity_key=key, name=name, **values)
        db.add(row); db.flush()
    return row


def source_summary(row: Evidence) -> dict:
    """row为存在的Evidence；返回可公开来源信息，不暴露本地路径、原文或内部凭据。"""
    meta = row.source_meta or {}
    return {'id': row.id, 'url': row.url, 'title': row.title, 'channel': row.channel,
            'published_at': row.published_at, 'captured_at': row.captured_at,
            'origin': meta.get('origin', 'agent'), 'reference': meta.get('source_id', ''),
            'evidence_kind': meta.get('evidence_kind', row.kind)}


def fact_view(row: CompanyFact) -> dict:
    """row为已保存事实；保留未知值与日期，争议/遮挡联系不给动作链接，不证明号码接通。"""
    data = serialize(row)
    data['numeric_value'] = float(row.numeric_value) if row.numeric_value is not None else None
    data.pop('provenance', None)
    period = row.report_period or ''
    meta = (row.provenance or {}).get('metadata_quote', '')
    # 明确起止区间才识别年度/半年；单个截至6月的时点不得推断半年收入。
    span = re.search(r'(\d{4})-01-01\s*(?:至|到|[~～—–])\s*\1-(12-31|06-30)', period)
    data['period_kind'] = ('quarter' if re.search('季度', period+meta) else 'half_year' if re.search('半年|1-6|1—6', period+meta)
        else 'annual' if re.search('年度|全年', period+meta) else 'annual' if span and span.group(2) == '12-31'
        else 'half_year' if span and span.group(2) == '06-30' else 'as_of')
    data['contact_href'] = ''
    value = str(row.raw_value or '')
    if row.status not in {'disputed', 'inactive', 'pending', 'masked'}:
        if row.field == 'phone': data['contact_href'] = phone_info(value)['href']
        elif row.field == 'email' and re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', value): data['contact_href'] = 'mailto:' + value
        elif row.field == 'website': data['contact_href'] = website_link(value)
    return data


def scope_ids(db, company: Company) -> list[str]:
    """db为会话、company为存在主体；集团视图仅扩展有证据且有持久ID的子公司，单体只含自身。"""
    result = {company.id}
    if company.view_scope == 'group':
        frontier = [company.id]
        for _ in range(4):
            children = set(db.scalars(select(CompanyRelation.counterparty_id).where(
                CompanyRelation.company_id.in_(frontier), CompanyRelation.relation_type == 'subsidiary',
                CompanyRelation.counterparty_id.is_not(None)))) - result
            if not children: break
            result.update(children); frontier = list(children)
    return sorted(result)


def summary_view(db, company: Company, facts=None) -> dict:
    """db为当前读取会话、company必存在、facts为空则查询；返回真实摘要、范围及统计，未知指标为null。"""
    facts = facts if facts is not None else list(db.scalars(select(CompanyFact).where(CompanyFact.company_id == company.id)))
    rows = sorted(facts, key=lambda x: (x.as_of or x.report_period or '', x.updated_at or ''), reverse=True)
    selected = {}
    for field in ('industry', 'business', 'products', 'headcount', 'revenue', 'address', 'website', 'stock_code'):
        eligible = [r for r in rows if r.field == field and r.status not in {'pending', 'disputed', 'inactive', 'masked'}]
        if field == 'revenue':
            eligible.sort(key=lambda r: (fact_view(r)['period_kind']=='annual', r.as_of), reverse=True)
        if field == 'products': eligible.sort(key=lambda r: (r.label == '报告产品系列', r.as_of), reverse=True)
        selected[field] = fact_view(eligible[0]) if eligible else None
    ids = scope_ids(db, company)
    projects = list(db.scalars(select(ProcurementProject).where(ProcurementProject.buyer_id.in_(ids),ProcurementProject.payload['duplicate_of'].as_string().is_(None))))
    purchases = [p for p in projects if p.record_kind == 'procurement']
    notice_count = db.scalar(select(func.count()).select_from(ProcurementNotice).where(ProcurementNotice.project_id.in_([p.id for p in purchases]),ProcurementNotice.payload['duplicate_of'].as_string().is_(None))) if purchases else 0
    buyer_names = list(db.scalars(select(Company.name).where(Company.id.in_({p.buyer_id for p in purchases})).order_by(Company.name))) if purchases else []
    relations = list(db.scalars(select(CompanyRelation).where(CompanyRelation.company_id == company.id)))
    periods = {kind:max((r.report_period for r in relations if r.relation_type == kind), default='') for kind in ('supplier','customer')}
    source_ids = set(s for f in facts for s in f.source_ids or [])
    source_ids.update(s for p in projects for s in p.source_ids or [])
    source_ids.update(s for r in relations for s in r.source_ids or [])
    source_ids.update(tender_source_ids(db,company.id))
    from .tender_service import tender_counts
    from .logo_service import logo_url
    return {**serialize(company), **tender_counts(db, [company.id])[company.id], 'logo_url': logo_url(company), 'facts': selected, 'fact_count': len(facts), 'source_count': len(source_ids),
            'project_count': sum(p.record_kind == 'procurement' for p in projects),
            'notice_count': notice_count, 'buyer_names': buyer_names,
            'procurement_categories': list(dict.fromkeys(p.payload.get('group') or p.payload.get('category') for p in purchases if p.payload.get('group') or p.payload.get('category')))[:3],
            'plan_count': sum(p.record_kind == 'plan' for p in projects),
            'supplier_count': len({r.counterparty_name for r in relations if r.relation_type == 'supplier' and r.report_period == periods['supplier']}),
            'customer_count': len({r.counterparty_name for r in relations if r.relation_type == 'customer' and r.report_period == periods['customer']})}


def list_companies(q='', entity_kind='', conversation_id='', page=1, page_size=20, visibility='primary') -> dict:
    """各筛选空串表示不限、page>=1/page_size<=100；返回数据库分页，坏会话404，不把关联主体套用主结果条件。"""
    with Session() as db:
        statement = select(Company)
        if q: statement = statement.where(Company.name.contains(q) | Company.short_name.contains(q))
        if entity_kind: statement = statement.where(Company.entity_kind == entity_kind)
        if visibility: statement = statement.where(Company.visibility == visibility)
        if conversation_id:
            if not db.get(Conversation, conversation_id): raise ServiceError('对话不存在', 404)
            statement = statement.where(Company.id.in_(select(ResearchCompany.company_id).where(ResearchCompany.conversation_id == conversation_id)))
        total = db.scalar(select(func.count()).select_from(statement.subquery())) or 0
        # 在分页前排序，最近核验/保存的企业优先；同秒用创建时间和ID保持稳定。
        rows = list(db.scalars(statement.order_by(Company.updated_at.desc(), Company.created_at.desc(), Company.id).offset((page - 1) * page_size).limit(page_size)))
        facts = list(db.scalars(select(CompanyFact).where(CompanyFact.company_id.in_([r.id for r in rows])))) if rows else []
        return {'items': [summary_view(db, r, [f for f in facts if f.company_id == r.id]) for r in rows],
                'total': total, 'page': page, 'page_size': page_size}


def company_detail(company_id: str) -> dict:
    """company_id非空；返回企业综合页数据与来源，缺失404，空采购不解释为实际采购为零。"""
    with Session() as db:
        company = db.get(Company, company_id)
        if not company: raise ServiceError('企业不存在', 404)
        facts = list(db.scalars(select(CompanyFact).where(CompanyFact.company_id == company_id).order_by(CompanyFact.as_of.desc(), CompanyFact.label)))
        ids = scope_ids(db, company)
        projects = list(db.scalars(select(ProcurementProject).where(or_(ProcurementProject.buyer_id.in_(ids),
            ProcurementProject.payload['research_company_id'].as_string() == company_id),ProcurementProject.payload['duplicate_of'].as_string().is_(None)).order_by(ProcurementProject.published_at.desc(), ProcurementProject.title)))
        notices = list(db.scalars(select(ProcurementNotice).where(ProcurementNotice.project_id.in_([p.id for p in projects]),ProcurementNotice.payload['duplicate_of'].as_string().is_(None)).order_by(ProcurementNotice.published_at.desc()))) if projects else []
        relations = list(db.scalars(select(CompanyRelation).where(CompanyRelation.company_id == company_id).order_by(CompanyRelation.relation_type, CompanyRelation.rank, CompanyRelation.counterparty_name)))
        source_ids = {s for r in [*facts, *projects, *notices, *relations] for s in r.source_ids or []}
        source_ids.update(tender_source_ids(db,company_id))
        sources = [source_summary(s) for s in db.scalars(select(Evidence).where(Evidence.id.in_(source_ids)).order_by(Evidence.published_at.desc()))]
        refs = entry_references(db, {'fact': [f.id for f in facts], 'project': [p.id for p in projects],
            'notice': [n.id for n in notices], 'relation': [r.id for r in relations]})
        fact_values = [{**fact_view(f), 'evidence_refs': refs.get(f.id, [])} for f in facts]
        grouped = {}
        for fact in fact_values: grouped.setdefault(fact['category'] or '其他公开资料', []).append(fact)
        project_values = []
        for project in projects:
            value = serialize(project)
            value['evidence_refs'] = refs.get(project.id, [])
            value['buyer_name'] = db.get(Company, project.buyer_id).name
            value['notices'] = [{**serialize(n), 'evidence_refs': refs.get(n.id, [])} for n in notices if n.project_id == project.id]
            value['temporal_status'] = ('截止已过 · 结果未知' if project.deadline and project.deadline[:10] < now()[:10] and not re.search('成交|已中标|终止', project.stage)
                                        else '候选公示 · 非最终成交' if '候选' in project.stage else project.stage or '阶段未核实')
            project_values.append(value)
        relation_values = []
        relation_keys = {}
        from .company_validation import reversed_subsidiary
        # 历史保存的方向相反子公司关系（本企业实为对方下属）不展示。
        relations = [r for r in relations if not (r.relation_type == 'subsidiary' and reversed_subsidiary(company.name, r.counterparty_name, r.quote))]
        for relation in relations:
            from .relation_service import relation_projection
            value = relation_projection(relation, company)
            value['evidence_refs'] = refs.get(relation.id, [])
            value['amount'] = float(relation.amount) if relation.amount is not None else None
            value['share'] = float(relation.share) if relation.share is not None else None
            key = stable_key(value['counterparty_name'], value['relation_type'], value['report_period'],
                value['amount'], value['share'], value['unit'], value['currency'])
            if key in relation_keys:
                previous = relation_keys[key]
                previous['source_ids'] = list(dict.fromkeys([*previous['source_ids'], *value['source_ids']]))
                previous['evidence_refs'] = list({r['binding_id']: r for r in [*previous['evidence_refs'], *value['evidence_refs']]}.values())
                continue
            relation_keys[key] = value
            relation_values.append(value)
        relation_summary = {}
        for kind in ('supplier', 'customer'):
            rows = [r for r in relation_values if r['relation_type'] == kind]
            period = max((r['report_period'] for r in rows), default='')
            current = [r for r in rows if r['report_period'] == period]
            units = {(r['unit'], r['currency']) for r in current}
            # 年份或币种不同不可直接加总；所有金额未知也不能展示“0元”。
            relation_summary[kind] = {'count':len(current), 'period':period,
                'amount':float(sum((Decimal(str(r['amount'])) for r in current), Decimal(0))) if current and len(units) == 1 and all(r['amount'] is not None for r in current) else None,
                'share':float(sum((Decimal(str(r['share'])) for r in current), Decimal(0))) if current and all(r['share'] is not None for r in current) else None,
                'unit':current[0]['unit'] if current and len(units) == 1 else '',
                'currency':current[0]['currency'] if current and len(units) == 1 else None}
        members = list(db.scalars(select(ResearchCompany).where(ResearchCompany.company_id == company_id).order_by(ResearchCompany.created_at.desc())))
        # 专项匹配回答保持独立于企业事实；只展示本企业已完成的真实模型任务，不借其他会话结论。
        analysis_job = db.scalar(select(Job).where(
            Job.metrics['target_company_id'].as_string() == company_id,
            Job.metrics['company_match_context'].as_boolean() == True,
            Job.status.in_(['completed', 'partial']), Job.reply_message_id.is_not(None)
        ).order_by(Job.finished_at.desc(), Job.created_at.desc()).limit(1))
        reply = db.get(Message, analysis_job.reply_message_id) if analysis_job else None
        analysis = {'job_id': analysis_job.id, 'conversation_id': analysis_job.conversation_id,
                    'created_at': analysis_job.finished_at or analysis_job.created_at,
                    'source_as_of': (analysis_job.metrics or {}).get('company_match_as_of', ''),
                    'status': analysis_job.status, 'origin': 'model_document_match',
                    'content': reply.content} if reply and reply.content else None
        header = summary_view(db, company, facts)
        for value in header['facts'].values():
            if value: value['evidence_refs'] = refs.get(value['id'], [])
        return {'header': header, 'fact_groups': grouped,
                'facts': fact_values, 'projects': project_values, 'relations': relation_values,
                'sources': sources, 'research_memberships': [serialize(m) for m in members],
                'coverage': (company.metadata_json or {}).get('coverage', []),
                'gaps': (company.metadata_json or {}).get('gaps', []),
                'procurement_summary': {'projects': sum(p.record_kind == 'procurement' for p in projects),
                    'notices': len(notices), 'plans': sum(p.record_kind == 'plan' for p in projects),
                    'quantity_disclosed': sum(p.payload.get('quantity') is not None for p in projects if p.record_kind == 'procurement')},
                'relation_summary': relation_summary, 'purchase_analysis': analysis,
                'tenders': {**tender_view(db, company_id),'reused':bool((company.metadata_json or {}).get('tender_reused')),
                            'diagnostics':(company.metadata_json or {}).get('tender_diagnostics',{})},
                'registration': registration_view(company)}


REGISTRATION_ORDER = ['法定代表人', '经营状态', '成立日期', '注册资本', '实缴资本', '工商注册号', '核准日期', '统一信用代码', '企业类型', '行业', '登记机关', '营业期限', '注册地址', '官网', '电话', '邮箱', '经营范围']


def registration_view(company: Company) -> dict | None:
    """company为主体；返回招投标平台公开的工商摘要（固定顺序）与来源链接，没有时返回None。"""
    info = dict((company.metadata_json or {}).get('registration') or {})
    # 通过逐字段证据门禁的工商事实直接投影到现有概要；不依赖旧聚合器私有metadata路径。
    with Session() as db:
        facts=list(db.scalars(select(CompanyFact).where(CompanyFact.company_id==company.id,CompanyFact.status=='verified').order_by(CompanyFact.updated_at.desc())))
        for fact in facts:
            if fact.label not in REGISTRATION_ORDER or fact.label in info:continue
            info[fact.label]=str(fact.raw_value)
            sources=[db.get(Evidence,sid) for sid in fact.source_ids]
            sources=[source for source in sources if source]
            if sources:
                info.setdefault('as_of',sources[0].captured_at)
                if sources[0].url.startswith(('http://','https://')):info.setdefault('source_url',sources[0].url)
    fields = [{'label': key, 'value': info[key]} for key in REGISTRATION_ORDER if info.get(key)]
    if not fields: return None
    return {'fields': fields, 'source_url': info.get('source_url', ''), 'sources': info.get('sources', []), 'as_of': info.get('as_of', '')}


def company_match_context(company_id: str) -> str:
    """读取指定企业的有界匹配上下文，不把模型判断写成企业事实。

    参数：company_id为持久企业ID且非空；空值或不存在抛ServiceError（404）。
    返回：最多22000字符的完整JSON，区分公开事实、项目/计划、来源和未读范围。
    异常：数据库读取错误向上传播，不用其他企业或历史对话替代目标资料。
    """
    detail = company_detail(company_id)
    facts = [row for row in detail['facts'] if row['status'] == 'verified'
             and row['field'] not in {'phone', 'email', 'fax', 'credit_code'}]
    # 采购依据优先；不用联系号码占用产品匹配的模型上下文，不改变公开原值。
    facts.sort(key=lambda row: (not bool(re.search('采购', row['category'] + row['label'])), row['field'] not in {'business', 'products', 'industry'}))
    projects = [row for row in detail['projects'] if row['record_kind'] in {'procurement', 'plan'}]
    payload = {'company_id': company_id, 'company_name': detail['header']['name'],
        'origin': detail['header']['origin'], 'research_as_of': detail['header']['research_as_of'],
        'facts': [{'label': row['label'], 'category': row['category'], 'note': row['note'][:600], 'value': str(row['raw_value'])[:600], 'unit': row['raw_unit'],
                   'period': row['report_period'] or row['as_of'], 'scope': row['scope'],
                   'source_ids': row['source_ids'], 'quote': row['quote'][:600],
                   'excerpted': len(str(row['raw_value'])) > 600 or len(row['quote']) > 600} for row in facts[:24]],
        'projects': [{'title': row['title'], 'buyer_name': row['buyer_name'], 'record_kind': row['record_kind'],
                      'project_number':row['project_number'],'published_at':row['published_at'],'deadline':row['deadline'],
                      'stage': row['temporal_status'], 'budget': row['payload'].get('budget'),
                      'unit': row['payload'].get('unit'), 'quantity': row['payload'].get('quantity'),
                      'specifications':str(row['payload'].get('specifications') or '')[:2000],
                      'quote':row['quote'][:1200],
                      'note': str(row['payload'].get('note') or '')[:600], 'source_ids': row['source_ids'],
                      'evidence_refs':[{k:ref.get(k) for k in ['source_id','role','quote']} for ref in row.get('evidence_refs',[])[:6]],
                      'notices':[{'title':n['title'],'stage':n['notice_type'],'published_at':n['published_at'],'deadline':n['deadline'],
                          'source_ids':n['source_ids']} for n in row.get('notices',[])[:4]],
                      'notice_count':len(row.get('notices',[])),
                      'excerpted':len(row['quote'])>1200 or len(str(row['payload'].get('specifications') or ''))>2000 or len(row.get('notices',[]))>4} for row in projects[:12]],
        'sources': [],
        'limitations': ['公开资料及历史计划不证明当前购买意向；已购状态、剩余预算和具体时间可能未知。'],
        'coverage': {'total_facts': len(facts), 'total_projects': len(projects), 'selected_facts': 0, 'selected_projects': 0}}
    available={row['id']: {'id':row['id'],'title':row['title'],'url':row['url']} for row in detail['sources']}
    def bind_sources():
        refs=set()
        for row in payload['facts']+payload['projects']:
            refs.update(row['source_ids'])
            refs.update(ref.get('source_id') for ref in row.get('evidence_refs',[]) if ref.get('source_id'))
            for notice in row.get('notices',[]):refs.update(notice['source_ids'])
        payload['sources']=[available[sid] for sid in sorted(refs) if sid in available]
        payload['coverage']['missing_source_ids']=sorted(refs-set(available))
    bind_sources()
    while len(json.dumps(payload, ensure_ascii=False)) > 21900:
        if payload['facts']: payload['facts'].pop()
        elif payload['projects']: payload['projects'].pop()
        else: raise ServiceError('企业匹配上下文超过限制')
        bind_sources()
    payload['coverage'].update(selected_facts=len(payload['facts']), selected_projects=len(payload['projects']))
    if payload['coverage']['missing_source_ids']:payload['limitations'].append('部分引用来源归档缺失，仅能使用已提供片段，原链接待核验。')
    return json.dumps(payload, ensure_ascii=False)


def company_research(company_id: str, values: dict) -> dict:
    """company_id必存在，values含client_token及可空conversation_id/content；创建真实研究任务，冲突409。"""
    detail = company_detail(company_id)
    conversation_id = values.get('conversation_id')
    with Session() as db:
        existing = db.scalar(select(Job).where(Job.client_token == values['client_token']))
        if existing:
            if (existing.metrics or {}).get('target_company_id') != company_id or conversation_id and existing.conversation_id != conversation_id:
                raise ServiceError('请求标识已用于其他研究', 409)
            if values.get('_operation') and any((existing.metrics or {}).get(k) != v for k, v in values['_operation'].items()):
                raise ServiceError('请求标识已用于不同操作', 409)
            return serialize(existing)
    if not conversation_id:
        with Session.begin() as db:
            conversation_id = stable_key('company-research', company_id, values['client_token'])
            if not db.get(Conversation, conversation_id):
                db.add(Conversation(id=conversation_id, title='研究 ' + detail['header']['name']))
    content = (values.get('content') or '继续研究该企业的客户画像、采购需求和上下游，核验缺失信息并保留来源').strip()
    return submit(conversation_id, {'content': content, 'client_token': values['client_token'],
        'attachment_ids': [], 'query_mode': 'company_query', 'company_id': company_id}, operation=values.get('_operation'))


def export_company(company_id: str):
    """company_id必存在；返回原值、口径及来源的XLSX字节流，不导出磁盘路径或凭据，缺失404。"""
    from openpyxl import Workbook
    detail = company_detail(company_id)
    workbook = Workbook(); sheet = workbook.active; sheet.title = '企业事实'
    sheet.append(['企业', '类别', '字段', '原值', '单位', '报告期', '主体口径', '核验状态', '引文', '来源ID', '备注'])
    for f in detail['facts']:
        sheet.append([detail['header']['name'], f['category'], f['label'], str(f['raw_value']), f['raw_unit'], f['report_period'] or f['as_of'], f['scope'], f['status'], f['quote'], ','.join(f['source_ids']), f['note']])
    project_sheet = workbook.create_sheet('采购项目'); project_sheet.append(['标题', '实际买方', '类型', '项目号', '阶段', '发布日期', '截止', '公开数据', '来源ID'])
    for p in detail['projects']:
        project_sheet.append([p['title'], p['buyer_name'], p['record_kind'], p['project_number'], p['temporal_status'], p['published_at'], p['deadline'], json.dumps(p['payload'], ensure_ascii=False), ','.join(p['source_ids'])])
    relation_sheet = workbook.create_sheet('上下游关系'); relation_sheet.append(['对手方', '关系', '年度', '年度交易额', '币种', '单位', '占比', '口径', '原文', '来源ID'])
    for r in detail['relations']:
        relation_sheet.append([r['counterparty_name'], r['relation_type'], r['report_period'], r['amount'], r['currency'], r['unit'], r['share'], r['scope'], r['quote'], ','.join(r['source_ids'])])
    source_sheet = workbook.create_sheet('来源'); source_sheet.append(['ID', '标题', 'URL', '发布日期', '采集时间', '来源类型'])
    for source in detail['sources']: source_sheet.append([source['id'], source['title'], source['url'], source['published_at'], source['captured_at'], source['origin']])
    for tab in workbook:
        tab.freeze_panes = 'A2'; tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            font = copy(cell.font); font.bold = True; cell.font = font
        for column in tab.columns: tab.column_dimensions[column[0].column_letter].width = min(55, max(15, max(len(str(c.value or '')) for c in column[:30]) + 2))
        # 公开文本不可成为Excel公式；保护原义而不改变数据库原值。
        for row in tab.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')): cell.data_type = 's'
    output = io.BytesIO(); workbook.save(output); output.seek(0); return output
