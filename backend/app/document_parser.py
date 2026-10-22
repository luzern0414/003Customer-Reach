"""本地文档解析：保留页码、表格、单元格与公式，不用模型补齐内容。"""
import io
import os
import shutil
import subprocess
import threading
import time
import zipfile
from datetime import date, datetime
from pathlib import Path
from pydantic import BaseModel, Field
from .config import ROOT
from .db import identifier
from .service import ServiceError

ALLOWED = {'.txt', '.csv', '.pdf', '.doc', '.docx', '.xls', '.xlsx'}
_OCR = None
_OCR_LOCK = threading.Lock()


class ParsedDocument(BaseModel):
    """确定性解析结果及可公开元信息；text只作为用户第一方产品资料。"""
    text: str
    truncated: bool = False
    metadata: dict = Field(default_factory=dict)


def _workspace() -> Path:
    directory = ROOT / '.codex-tmp/documents' / identifier()
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _executable(kind: str) -> str:
    configured = os.getenv('CLIENTREACH_' + kind.upper(), '')
    if configured and Path(configured).is_file():
        return configured
    found = shutil.which('soffice.com' if kind == 'soffice' else 'pdftoppm')
    if found:
        return found
    candidates = [Path('C:/Program Files/LibreOffice/program/soffice.com')] if kind == 'soffice' else list(
        (Path.home() / '.cache/codex-runtimes').glob('*/dependencies/native/poppler/Library/bin/pdftoppm.exe'))
    match = next((x for x in candidates if x.is_file()), None)
    if not match:
        raise ServiceError('解析依赖缺失：请配置' + ('LibreOffice' if kind == 'soffice' else 'Poppler'))
    return str(match)


def _ocr(content: bytes | Path) -> tuple[str, float | None]:
    global _OCR
    # 实例会更新参数，初始化和推理共用锁，避免并发上传互相修改配置。
    with _OCR_LOCK:
        if _OCR is None:
            from rapidocr import RapidOCR
            _OCR = RapidOCR(params={'Global.model_root_dir': str(ROOT / 'data/ocr-models'),
                'Global.log_level': 'warning', 'EngineConfig.onnxruntime.intra_op_num_threads': 2,
                'EngineConfig.onnxruntime.inter_op_num_threads': 1})
        result = _OCR(content)
    words, scores = result.txts or (), result.scores or ()
    return '\n'.join(words), round(sum(scores) / len(scores), 3) if scores else None


def _pdf(content: bytes, meta: dict) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ServiceError('加密PDF需要先解除密码后上传')
    meta.update(method='PDF文本 + 按页OCR', pages=len(reader.pages), ocr_pages=[], page_details=[])
    parts, work = [], None
    for number, page in enumerate(reader.pages[:100], 1):
        value = page.extract_text() or ''
        detail = {'page': number, 'method': 'text', 'characters': len(value)}
        # 页面含栅格图片时仍做OCR，标题文字层不能证明扫描正文已经提取。
        has_images=bool(page.images)
        if not value.strip() or has_images:
            if len(meta['ocr_pages']) >= 10:
                meta['warnings'].append(f'第{number}页未识别：达到10页OCR上限')
                meta['limited'] = True
                if value.strip():
                    detail['method']='text · 图片未识别'
                    meta['page_details'].append(detail)
                    parts.append(f'[第{number}页 · {detail["method"]}]\n{value.strip()}')
                continue
            if work is None:
                work = _workspace()
                (work / 'input.pdf').write_bytes(content)
            prefix = work / f'page-{number}'
            subprocess.run([_executable('pdftoppm'), '-f', str(number), '-l', str(number),
                '-scale-to', '2000', '-singlefile', '-png', str(work / 'input.pdf'), str(prefix)],
                check=True, timeout=60, capture_output=True)
            scanned, confidence = _ocr(prefix.with_suffix('.png'))
            # 保留准确文字层并标记OCR补充，避免用识别误差覆盖原始型号/金额。
            native=value.strip()
            value='\n'.join([native,'[图片与页面OCR补充]',scanned.strip()]) if native and scanned.strip() else native or scanned
            meta['ocr_pages'].append(number)
            detail.update(method='text+ocr' if native else 'ocr', confidence=confidence, characters=len(value))
            if not value.strip():
                meta['warnings'].append(f'第{number}页未识别到文字')
                meta['limited'] = True
            elif native and has_images and not scanned.strip():
                meta['warnings'].append(f'第{number}页图片未识别到文字，仅保留文字层')
                meta['limited'] = True
            elif confidence is not None and confidence < .85:
                meta['warnings'].append(f'第{number}页OCR置信度偏低，请核对型号和数字')
        meta['page_details'].append(detail)
        if value.strip():
            parts.append(f'[第{number}页 · {detail["method"]}]\n{value.strip()}')
    if len(reader.pages) > 100:
        meta['limited'] = True
        meta['warnings'].append('只读取前100页')
    if meta['ocr_pages']:
        meta['warnings'].append('扫描页由OCR识别，未恢复完整表格布局；型号和金额请核对原件')
    return '\n\n'.join(parts)


def _docx(content: bytes, meta: dict) -> str:
    from docx import Document
    from docx.table import Table
    doc = Document(io.BytesIO(content))
    meta.update(method='Word段落/表格/页眉页脚', tables=0, images_ocr=0)

    def blocks(parent, depth=0):
        if depth > 12:
            meta['limited'] = True
            meta['warnings'].append('嵌套表格超过12层，超出部分未读取')
            return []
        out = []
        for block in parent.iter_inner_content():
            if isinstance(block, Table):
                meta['tables'] += 1
                table_rows=[]
                for index, row in enumerate(block.rows, 1):
                    seen, cells = set(), []
                    for cell in row.cells:
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        cells.append(' / '.join(blocks(cell, depth + 1)))
                    if any(cells):
                        table_rows.append(f'行{index}: ' + ' | '.join(cells))
                # 表格结构不代表有效正文；空表不能靠人工标记绕过最终空白检查。
                if table_rows:out.extend(['[表格]',*table_rows])
            elif block.text.strip():
                out.append(block.text)
        return out

    parts = blocks(doc)
    seen_headers = set()
    for section in doc.sections:
        for label, header in [('页眉', section.header), ('页脚', section.footer),
            ('首页页眉', section.first_page_header), ('首页页脚', section.first_page_footer),
            ('偶数页页眉', section.even_page_header), ('偶数页页脚', section.even_page_footer)]:
            if header.part.partname in seen_headers:
                continue
            seen_headers.add(header.part.partname)
            values = blocks(header)
            if values:
                parts.append(f'[{label}]\n' + '\n'.join(values))
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        images = [x for x in archive.namelist() if x.startswith('word/media/')]
        for item in images[:10]:
            try:
                value, confidence = _ocr(archive.read(item))
                meta['images_ocr'] += 1
                if value:
                    parts.append(f'[插图OCR · {Path(item).name}]\n{value}')
                else:
                    meta['warnings'].append(f'插图{Path(item).name}没有识别到文字')
                if confidence is not None and confidence < .85:
                    meta['warnings'].append(f'插图{Path(item).name}识别置信度偏低')
            except Exception:
                meta['limited'] = True
                meta['warnings'].append(f'插图{Path(item).name}未解析，可能是不支持的矢量图片')
        if len(images) > 10:
            meta['limited'] = True
            meta['warnings'].append('只识别前10张插图')
        xml = archive.read('word/document.xml')
        if any(marker in xml for marker in [b'txbxContent', b'<w:del ', b'<w:ins ', b'footnoteReference', b'endnoteReference']):
            meta['limited'] = True
            meta['warnings'].append('包含文本框、修订或脚注，这些特殊结构未完整提取')
    if images:
        meta['warnings'].append('插图OCR列在正文后，不表示原文位置；识别结果需核对原件')
    return '\n'.join(parts)


def _doc(content: bytes, meta: dict) -> str:
    work = _workspace()
    source = work / 'input.doc'
    source.write_bytes(content)
    profile = work / 'lo-profile'
    (profile / 'user').mkdir(parents=True)
    # 独立目录与高宏安全级别，避免使用用户正在编辑的Office实例。
    (profile / 'user/registrymodifications.xcu').write_text(
        '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" '
        'oor:op="fuse"><value>3</value></prop></item></oor:items>', encoding='utf-8')
    subprocess.run([_executable('soffice'), '-env:UserInstallation=' + profile.as_uri(), '--headless',
        '--convert-to', 'docx:Office Open XML Text', '--outdir', str(work), str(source)],
        capture_output=True, check=True, timeout=90)
    converted = work / 'input.docx'
    if not converted.is_file():
        raise ServiceError('Word旧格式转换失败，请另存为DOCX后上传')
    content = converted.read_bytes()
    _zip_guard(content)
    result = _docx(content, meta)
    meta['method'] = 'LibreOffice DOC转换 + ' + meta['method']
    meta['warnings'].append('DOC经LibreOffice转换，复杂排版可能变化')
    return result


def _value(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d %H:%M:%S') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    return str(value) if value is not None else ''


def _xlsx(content: bytes, meta: dict) -> str:
    from openpyxl import load_workbook
    formulas = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    values = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    meta.update(method='Excel单元格/公式/缓存值', sheets=[], formula_without_cache=0)
    parts = []
    try:
        for sheet in formulas:
            output, count = [], 0
            if sheet.max_row > 1000 or sheet.max_column > 30:
                meta['limited'] = True
                meta['warnings'].append(f'{sheet.title}仅读取前1000行/30列')
            bounds = dict(max_row=min(sheet.max_row, 1000), max_col=min(sheet.max_column, 30))
            # 缓存簿按行同步迭代，避免read_only随机逐格访问反复扫描工作表。
            for row, cached_row in zip(sheet.iter_rows(**bounds), values[sheet.title].iter_rows(**bounds)):
                entries = []
                for cell, cached_cell in zip(row, cached_row):
                    if cell.value is None:
                        continue
                    value = _value(cell.value)
                    if cell.data_type == 'f':
                        cached = cached_cell.value
                        if cached is None:
                            meta['formula_without_cache'] += 1
                            value += ' [无缓存值，未计算]'
                        else:
                            value += f' [缓存值: {_value(cached)}]'
                    entries.append(f'{cell.coordinate}: {value}')
                    count += 1
                if entries:
                    output.append(' | '.join(entries))
            merged = [str(x) for x in sheet.merged_cells.ranges]
            meta['sheets'].append({'name': sheet.title, 'cells': count, 'merged_ranges': merged[:100]})
            if output:
                parts.append(f'[工作表: {sheet.title}]\n' + ('合并单元格: ' + ', '.join(merged[:100]) + '\n' if merged else '') + '\n'.join(output))
        if meta['formula_without_cache']:
            meta['warnings'].append(f'{meta["formula_without_cache"]}个公式没有缓存结果，保留公式但未自动计算')
        if formulas._external_links:
            meta['warnings'].append('外部链接未刷新，缓存值不保证最新')
    finally:
        formulas.close()
        values.close()
    return '\n\n'.join(parts)


def _xls(content: bytes, meta: dict) -> str:
    import xlrd
    from openpyxl.utils import get_column_letter
    book = xlrd.open_workbook(file_contents=content)
    meta.update(method='BIFF Excel缓存单元格', sheets=[])
    meta['warnings'].append('XLS读取保存时的单元格值，不重新计算或提取公式表达式')
    parts = []
    for sheet in book.sheets():
        output, count = [], 0
        if sheet.nrows > 1000 or sheet.ncols > 30:
            meta['limited'] = True
            meta['warnings'].append(f'{sheet.name}仅读取前1000行/30列')
        for row in range(min(sheet.nrows, 1000)):
            entries = []
            for col in range(min(sheet.ncols, 30)):
                cell = sheet.cell(row, col)
                if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                    continue
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = xlrd.error_text_from_code.get(value, '#ERROR')
                entries.append(f'{get_column_letter(col+1)}{row+1}: {_value(value)}')
                count += 1
            if entries:
                output.append(' | '.join(entries))
        meta['sheets'].append({'name': sheet.name, 'cells': count})
        if output:
            parts.append(f'[工作表: {sheet.name}]\n' + '\n'.join(output))
    book.release_resources()
    return '\n\n'.join(parts)


def _zip_guard(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
            raise ServiceError('文档解压内容过大，请拆分文件')


def parse_document(suffix: str, content: bytes) -> ParsedDocument:
    """suffix须含点且小写，content非空<=10MB；返回文本/范围/警告，不补写未读内容。

    损坏、加密或空白抛400，不支持抛415，超限抛413。超过100页/10页OCR/
    1000行/30列/12万字不推测，返回truncated及警告。
    """
    if suffix not in ALLOWED:
        raise ServiceError('支持PDF、DOC/DOCX、XLS/XLSX、CSV和TXT产品文件', 415)
    if not content or len(content) > 10 * 1024 * 1024:
        raise ServiceError('附件不能为空且不能超过10MB', 413)
    started = time.perf_counter()
    meta = {'format': suffix[1:], 'warnings': [], 'limited': False}
    try:
        if suffix in {'.docx', '.xlsx'}:
            _zip_guard(content)
        if suffix in {'.txt', '.csv'}:
            try:
                result = content.decode('utf-8-sig')
                meta['encoding'] = 'UTF-8'
            except UnicodeDecodeError:
                result = content.decode('gb18030')
                meta['encoding'] = 'GB18030'
            meta['method'] = '文本解码'
        else:
            result = {'.pdf': _pdf, '.docx': _docx, '.doc': _doc, '.xlsx': _xlsx, '.xls': _xls}[suffix](content, meta)
    except ServiceError:
        raise
    except subprocess.TimeoutExpired:
        raise ServiceError('文档转换或扫描页识别超时，请拆分后上传')
    except Exception as exc:
        raise ServiceError('文件无法解析或内容格式损坏，请检查后重新上传') from exc
    if not result.strip():
        raise ServiceError('文件没有可分析的文字内容，扫描件请核对清晰度')
    if len(result) > 120000:
        meta['limited'] = True
        meta['warnings'].append('解析结果超过12万字符，仅保留前12万字符')
    meta['warnings'] = list(dict.fromkeys(meta['warnings']))
    meta['elapsed_ms'] = round((time.perf_counter() - started) * 1000)
    return ParsedDocument(text=result[:120000], truncated=meta.pop('limited'), metadata=meta)
