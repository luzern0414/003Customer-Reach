"""CRM业务服务：持久化、候选事实校验、任务事件与人工修改。"""
import hashlib
import json
import re
import time
from pathlib import Path
from sqlalchemy import select, func, true
from .db import Session, Conversation, Message, Job, JobEvent, Lead, Followup, Setting, Evidence, now, serialize
from .config import ROOT, credentials
from .schemas import SettingsInput, Candidate, LeadInput
from .observability import logger
from .identity import city_key,city_values

LOG = logger('agent')


def evidence_detail(source_id: str) -> dict:
    """source_id为非空持久来源ID；返回实际正文和公开元数据，磁盘路径不出API，缺失404。"""
    with Session() as db:
        row=db.get(Evidence,source_id)
        if not row:raise ServiceError('来源不存在',404)
        value=serialize(row)
        allowed={'origin','dataset','source_id','evidence_kind','format','sha256','hash_object','backend','opencli_version','truncated','requested_url','quote_kind'}
        value['source_meta']={key:item for key,item in (row.source_meta or {}).items() if key in allowed}
        return value


class ServiceError(Exception):
    """可向用户说明的业务错误，不包含provider原始响应。"""
    def __init__(self, message: str, status: int = 400):
        """message为可公开非空文案，status为HTTP错误码；返回业务异常对象。"""
        super().__init__(message)
        self.status = status


def settings() -> dict:
    """返回非敏感设置与配置状态；读取失败只显示未配置，不暴露密钥。"""
    with Session() as db:
        row = db.get(Setting, 'app')
        result = {**SettingsInput().model_dump(), **(row.value if row else {})}
    try:
        url, _ = credentials()
        result.update(model_configured=True, base_url=url)
    except ValueError:
        result.update(model_configured=False, base_url='')
    return {**result, 'search_provider': 'Exa MCP', 'database': 'SQLite',
            'credential_source': 'config/apikey.txt', 'version': '0.3.0', 'agent_engine':'LangGraph'}


def save_settings(values: dict) -> dict:
    """保存已校验的非敏感values；不可为空，不接受密钥；返回最新设置。"""
    values = SettingsInput(**values).model_dump()
    with Session.begin() as db:
        row = db.get(Setting, 'app')
        if row:
            row.value = values
        else:
            db.add(Setting(key='app', value=values))
    return settings()


def conversation_detail(conversation_id: str) -> dict:
    """读取非空对话ID及消息/任务；不存在抛404，返回按时间排序的持久化记录。"""
    with Session() as db:
        row = db.get(Conversation, conversation_id)
        if not row:
            raise ServiceError('对话不存在', 404)
        from .db import Attachment, ToolCall
        from .files import attachment_meta
        from .tool_views import tool_view
        from .company_models import ResearchDataset
        from .research_dataset import dataset_summary
        messages = list(db.scalars(select(Message).where(Message.conversation_id == conversation_id).order_by(Message.id)))
        jobs = list(db.scalars(select(Job).where(Job.conversation_id == conversation_id).order_by(Job.created_at.desc(), Job.metrics['submitted_epoch'].as_float().desc(), Job.id.desc())))
        grouped = {x.id: [] for x in jobs}
        # 一次查询完整会话轨迹，历史轮次不逐个请求任务，避免N+1与页面加载变慢。
        for call in db.scalars(select(ToolCall).where(ToolCall.job_id.in_(grouped)).order_by(ToolCall.created_at, ToolCall.id)):
            if call.channel != 'model':
                grouped[call.job_id].append(tool_view(call))
        job_values = []
        datasets={key:[] for key in grouped}
        for dataset in db.scalars(select(ResearchDataset).where(ResearchDataset.job_id.in_(grouped)).order_by(ResearchDataset.created_at,ResearchDataset.id)):
            datasets[dataset.job_id].append(dataset)
        for item in jobs:
            value = serialize(item)
            if item.clock_at and item.status in {'running', 'cancelling'}:
                value['elapsed_ms'] = (value['elapsed_ms'] or 0) + max(0, round((time.time() - item.clock_at) * 1000))
            job_values.append({**value, 'tools': grouped[item.id],'research_summary':dataset_summary(datasets[item.id])})
        file_ids = {key for message in messages for key in message.attachment_ids or []}
        files = [attachment_meta(x) for x in db.scalars(select(Attachment).where(Attachment.id.in_(file_ids)))]
        from .conversation import short_title
        return {**serialize(row), 'title':short_title(row.title), 'messages': [serialize(x) for x in messages], 'jobs': job_values, 'attachments': files}


def submit(conversation_id: str, values: dict, operation: dict | None = None) -> dict:
    """提交任务；conversation_id非空、values为校验参数，operation为内部操作元数据（空沿普通流程）。

    返回持久任务；重复token操作范围不同/已有活动任务409，无资源404，不接收客户端任意metrics。
    """
    from .db import Attachment
    with Session.begin() as db:
        existing = db.scalar(select(Job).where(Job.client_token == values['client_token']))
        if existing:
            if existing.conversation_id != conversation_id:
                raise ServiceError('请求标识已用于其他对话', 409)
            if operation and any((existing.metrics or {}).get(key) != value for key, value in operation.items()):
                raise ServiceError('请求标识已用于不同操作', 409)
            return serialize(existing)
        conv = db.get(Conversation, conversation_id)
        if not conv:
            raise ServiceError('对话不存在', 404)
        active = db.scalar(select(Job).where(Job.conversation_id == conversation_id,
                                             Job.status.in_(['queued', 'running', 'cancelling'])))
        if active:
            raise ServiceError('这段对话的任务仍在执行，可先暂停或等待完成', 409)
        for file_id in values.get('attachment_ids', []):
            if not db.get(Attachment, file_id):
                raise ServiceError('附件不存在，请重新上传', 404)
        prompt = values['content'].strip()
        if not prompt:
            raise ServiceError('请输入产品或获客要求')
        from .conversation import short_title
        if conv.title == '新的获客对话':conv.title = short_title(prompt)
        company_id = values.get('company_id')
        if values.get('query_mode') == 'company_match' and (not company_id or not values.get('attachment_ids')):
            raise ServiceError('产品匹配需要指定企业和已上传资料')
        if company_id:
            from .company_models import Company
            if not db.get(Company, company_id): raise ServiceError('企业不存在', 404)
        conv.updated_at = now()
        db.add(Message(conversation_id=conversation_id, role='user', content=prompt,
                       attachment_ids=values.get('attachment_ids', [])))
        job = Job(conversation_id=conversation_id, client_token=values['client_token'],
                  prompt=prompt, attachment_ids=values.get('attachment_ids', []),
                  query_mode=values.get('query_mode','auto'),
                  metrics={'submitted_epoch':time.time(),'intent_pending':True,'pipeline_version':'research-v3','conversation_title':short_title(conv.title),
                           **({'target_company_id': company_id} if company_id else {}), **(operation or {})})
        db.add(job)
        db.flush()
        return serialize(job)


def resume_task(job_id: str) -> dict:
    """job_id须非空；显式恢复时建立重试轮次，不退还额度，限流冷却/活动冲突抛409，不存在404。"""
    from .db import ToolCall
    with Session.begin() as db:
        row=db.get(Job,job_id)
        if not row:raise ServiceError('任务不存在',404)
        if row.status not in {'paused','failed','partial'}:raise ServiceError('此任务当前不支持恢复',409)
        if (row.metrics or {}).get('pending_cleanup'):raise ServiceError('上一轮请求仍在停止清理，请稍后继续',409)
        if db.scalar(select(Job).where(Job.conversation_id==row.conversation_id,Job.status.in_(['queued','running','cancelling']))):
            raise ServiceError('此对话已有执行中的任务',409)
        # 单站点受限不封锁其他来源；供应商限流才挡住整轮公开检索恢复。
        deadline=db.scalar(select(func.max(ToolCall.retry_after_epoch)).where(
            ToolCall.result['error'].as_string()=='provider_rate_limited')) or 0
        if deadline>time.time():raise ServiceError(f'检索服务暂时受限，请{max(1,int(deadline-time.time())+1)}秒后重试',409)
        row.status,row.cancel_requested,row.error='queued',False,''
        row.metrics={**(row.metrics or {}),'resume_epoch':time.time()};row.updated_at=now()
        return serialize(row)


def record_event(job_id: str, stage: str, message: str, progress: int | None = None,
                 level: str = 'info') -> None:
    """记录任务进度；各文本为非敏感程序文案，progress为空保留原值，不存在任务忽略。"""
    with Session.begin() as db:
        job = db.get(Job, job_id)
        if not job:
            return
        job.stage, job.updated_at = stage, now()
        if progress is not None:
            job.progress = max(0, min(100, progress))
        db.add(JobEvent(job_id=job_id, stage=stage, message=message, level=level))
    LOG.log(30 if level == 'warning' else 20, 'job_event',
            extra={'details': {'job_id': job_id, 'stage': stage, 'progress': progress}})


def checked_candidate(candidate: Candidate, source_rows: dict) -> tuple[dict, list]:
    """校验candidate与实际工具source_rows；均不可为空，无原文/身份证据抛业务错误。

    返回可入库事实和引文列表。字段值必须出现在逐字引文中；统计和采购口径额外约束。
    """
    clean, citations = {}, []
    from .public_directory import branch_key
    name_branches={branch_key(source_rows[fact.source_id].get('url','')) for fact in candidate.facts
        if fact.field=='name' and fact.source_id in source_rows and branch_key(source_rows[fact.source_id].get('url',''))}
    if len(name_branches)>1:raise ServiceError('商户分店证据不一致，不能合并不同分店的联系方式')
    for fact in candidate.facts:
        source = source_rows.get(fact.source_id)
        if not source:
            raise ServiceError('引用的来源不是本任务实际读取的资料')
        branch=branch_key(source.get('url',''))
        if name_branches and branch and branch not in name_branches:
            raise ServiceError('商户分店证据不一致，不能合并不同分店的联系方式')
        normalize = lambda value: re.sub(r'\s+', '', value)
        from .industry_navigation import factual_body
        if normalize(fact.quote) not in normalize(factual_body(source['body'])):
            raise ServiceError(f'{fact.field}的引文未能在原始资料中找到')
        if normalize(fact.value) not in normalize(fact.quote):
            raise ServiceError(f'{fact.field}的值不是引文原样内容，请保留原文而非扩写')
        if fact.field=='name' and normalize(fact.value)!=normalize(candidate.display_name):
            # 每条name都代表本候选的完整主体，另一简称/品牌不能作为第二name混入争议账本。
            raise ServiceError('客户名称必须有对应原文name字段证据')
        if fact.field=='industry':
            # 逐字命中只证明文字存在，不能把地区或产品分类误填到行业字段。
            if normalize(fact.value).removesuffix('市')==normalize(candidate.city).removesuffix('市') or not re.search('行业|零售业|批发业|制造业|建筑业',fact.quote):
                raise ServiceError('行业需要明确行业分类原文，城市或主营品类不能替代行业；缺证据请省略industry')
        if fact.field=='categories' and re.search(r'市场(?:类型|类别)[\s|:：]*'+re.escape(normalize(fact.value)),normalize(fact.quote)) and re.search('市场|商场|卖场|家居城|灯饰城',fact.value):
            # 场所类型可证明经营方角色，但不是它经营的具体商品；不能用跨站补查污染已有主营字段。
            raise ServiceError('市场类型不能替代主营品类；缺证据请省略categories')
        if fact.field in {'revenue', 'headcount'}:
            from .statistics_evidence import valid_statistics
            if not valid_statistics(fact.scope,fact.effective_date,fact.quote):
                raise ServiceError('营收/人数需要统计日期和主体口径，否则应留空')
        if fact.field == 'revenue' and ('注册资本' in fact.quote or not re.search('营收|营业收入|销售收入', fact.quote)):
            raise ServiceError('注册资本不能替代营业收入')
        if fact.field == 'purchase_price' and (not re.search('成交|进货|供货价|采购单价', fact.quote) or re.search('预算|上限|零售', fact.quote)):
            raise ServiceError('零售价或预算上限不能写成实际采购价')
        if fact.field == 'supplier' and not re.search('供应商|供货|供方', fact.quote):
            raise ServiceError('入驻品牌或代理声明不能写成采购供应商')
        if fact.field=='address' and candidate.buyer_role=='retailer':
            from .public_directory import mall_listing,mall_address_quote
            if mall_listing(candidate.display_name,factual_body(source['body'])) and not mall_address_quote(candidate.display_name,fact.value,factual_body(source['body'])):
                raise ServiceError('商户地址必须来自该商户独立记录，不能继承商场总地址')
        if fact.field=='phone' and (re.search(r'[*＊…]',fact.value) or len(re.findall(r'\d',fact.value))<7):
            raise ServiceError('电话未完整公开，不能补全掩码或作为完整联系信息保存')
        if fact.field=='phone':
            from .contact_owner import wechat_position
            body=normalize(factual_body(source['body']))
            if wechat_position(fact.value,body) is not None and not re.search(r'(?:联系电话|座机电话|固定电话|手机号码|移动电话|电话|手机|tel|phone)[：:|]*(?:[-*+•])?'+re.escape(normalize(fact.value))+r'(?![A-Za-z0-9])',body,re.I):
                raise ServiceError('电话未明确公开，不能把微信号当作电话号码')
        if fact.field=='phone' and re.search(r'店\s+入驻品牌',source['body']):
            from .public_directory import mall_listing
            listing=mall_listing(candidate.display_name,source['body'])
            if not listing or listing['phone']!=fact.value:raise ServiceError('商场电话未关联该门店，不能借用其他商户号码')
        elif fact.field=='phone' and source.get('channel') not in {'procurement','ccgp','sz_procurement'}:
            from .contact_owner import public_phone
            if not public_phone(candidate.display_name,fact.value,factual_body(source['body']),source.get('title','')):
                raise ServiceError('电话未与该客户主体建立关联，不能借用平台或其他企业号码')
        if fact.field=='email' and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',fact.value):
            raise ServiceError('邮箱必须是原文中完整公开的地址')
        if fact.field=='wechat':
            from .contact_owner import public_wechat
            if not public_wechat(candidate.display_name,fact.value,fact.quote,factual_body(source['body']),source.get('title','')):
                raise ServiceError('微信必须明确公开并与该客户主体关联，不能由电话或掩码推断')
        if fact.field=='website':
            from .contact_view import website_link
            if not website_link(fact.value):raise ServiceError('网站地址必须完整且无嵌入凭据，不能补造域名')
        if fact.field=='contact_name':
            from .contact_owner import contact_identity
            if not contact_identity(fact.value,fact.quote):raise ServiceError('联系人必须有公开联系身份，不能把法定代表人当采购联系人')
        if fact.field in {'phone','email','wechat','contact_name','contact_position'} and source.get('channel') in {'procurement','ccgp','sz_procurement'}:
            from .contact_owner import buyer_contact
            if not buyer_contact(candidate.display_name,fact.value,source['body']):raise ServiceError('联系方式未能定位到该买方，不能借用代理或平台号码')
        value={'value': fact.value, 'source_id': fact.source_id,'scope': fact.scope, 'effective_date': fact.effective_date,
               'verification': '原文匹配；时效及实际业务待核'}
        previous=clean.get(fact.field)
        # 同一结构化批次的跨站冲突也必须保留，不能以“最后一个字段覆盖”悄悄选择电话或地址。
        if previous and previous['value']!=value['value']:
            alternatives=list(previous.get('alternatives',[]))
            from .contact_view import coexisting_phones
            parallel=fact.field=='phone' and coexisting_phones(candidate.display_name,previous,value,source_rows)
            if parallel:value={**value,'relation':'additional_contact'}
            if value not in alternatives:alternatives.append(value)
            clean[fact.field]={**previous,'verification':previous['verification'] if parallel else 'disputed','alternatives':alternatives}
        else:clean[fact.field]=previous if previous and previous.get('verification')=='disputed' else value
        citations.append({'field': fact.field, 'value': fact.value, 'quote': fact.quote,
                          'source_id': fact.source_id, 'url': source['url'],
                          'title': source['title'], 'captured_at': source['captured_at'], 'kind': source['kind'],'channel':source.get('channel','')})
    name = clean.get('name')
    if not name or re.sub(r'\s+', '', candidate.display_name)!=re.sub(r'\s+', '', name['value']):
        raise ServiceError('客户名称必须有对应原文name字段证据')
    if candidate.city and not any(candidate.city in x['quote'] for x in citations):
        from .public_directory import mall_city_quote
        for item in list(citations):
            if item['field']!='name':continue
            source=source_rows[item['source_id']]
            quote=mall_city_quote(candidate.display_name,candidate.city,source['body'])
            if quote:
                # 商户短名和城市介绍相隔多段，不拼接name引文；保留真实单分店介绍的独立城市证据。
                citations.append({**item,'field':'city','value':candidate.city,'quote':quote})
                break
    if candidate.city and not any(candidate.city in x['quote'] for x in citations):
        raise ServiceError('城市必须出现在该客户的引文中，不能从全国目录推断门店位置')
    if candidate.buyer_role!='unknown':
        signals={'dealer':'经销|代理|授权|批发|销售点','retailer':'门市|门店|零售|专卖店|专营店|直营门市|店铺|灯具店|灯饰店|入驻品牌',
                 'mall':'家居城|灯饰城|灯(?:饰|具)(?:专业)?市场|卖场|商场|灯饰商城','project_buyer':'采购人|采购方|招标人',
                 'manufacturer':'制造|生产|厂家|工厂'}
        role=[x for x in citations if x['field']=='buyer_role']
        if not role or not any(re.search(signals[candidate.buyer_role],x['value']) for x in role):
            raise ServiceError('买家角色缺少对应原文证据，应留为unknown继续补查')
        entity_role=False
        from .evidence_role import role_denied,competing_subject,role_location
        for item in role:
            if role_denied(item['value'],item['quote']):continue
            # 目录常将多主体用逗号串成一句；必须在同一主谓片段，不能仅按句号分隔。
            for clause in re.split(r'[。\n；;，,]',item['quote']):
                normalized=normalize(clause);entity=normalize(candidate.display_name)
                if entity in normalized and not role_denied(item['value'],clause) and not role_location(item['value'],normalized.replace(entity,'')) and not competing_subject(candidate.display_name,clause) and re.search(signals[candidate.buyer_role],normalized.replace(entity,'')) and normalize(item['value']) in normalized.replace(entity,''):
                    entity_role=True
        if not entity_role:
            from .evidence_role import scoped_role
            for item in role:
                source=source_rows[item['source_id']]
                same_name=normalize(candidate.display_name) in normalize(source.get('body','')) or any(x['field']=='name' and x['source_id']==item['source_id'] and normalize(candidate.display_name) in normalize(x['value']) for x in citations)
                if scoped_role(candidate.display_name,item['value'],item['quote'],source.get('title',''),same_name,source.get('body','')):entity_role=True
            if not entity_role:raise ServiceError('角色原文没有与该客户主体建立关联')
        if candidate.buyer_role=='dealer' and any('采购代理' in x['quote'] or '招标代理' in x['quote'] for x in role):
            raise ServiceError('采购代理机构不是产品经销商')
    from .knowledge import invalidated_fact
    for fact in candidate.facts:
        if invalidated_fact(candidate.display_name,fact,source_rows[fact.source_id]):
            raise ServiceError('该事实已被用户标记有误或失效，不能复用旧证据')
    return clean, citations


def save_candidate(candidate: Candidate, source_rows: dict, job_id: str) -> dict:
    """保存校验后的候选；job_id非空，按名称/城市/地址去重，不覆盖人工字段及跟进状态。"""
    clean, citations = checked_candidate(candidate, source_rows)
    address = clean.get('address', {}).get('value', '')
    from .public_directory import branch_key
    branch=next((branch_key(item.get('url','')) for item in citations if item['field']=='name' and branch_key(item.get('url',''))),'')
    identity=f'{candidate.display_name}|{city_key(candidate.city)}|{address}'+('|'+branch if branch else '')
    fingerprint = hashlib.sha256(re.sub(r'\s+', '', identity).encode()).hexdigest()
    with Session.begin() as db:
        if branch:
            archived=list(db.scalars(select(Lead).where(Lead.display_name==candidate.display_name,Lead.city.in_(city_values(candidate.city)),Lead.archived.is_(True))))
            if any(branch in {branch_key(item.get('url','')) for item in lead.evidence or [] if item.get('field')=='name'} for lead in archived):
                raise ServiceError('客户已归档，不能自动重新保存')
        if not branch:
            # 旧指纹可能使用带市原值；相同地址的归档主体不能借新比较键重新出现。
            archived=list(db.scalars(select(Lead).where(Lead.display_name==candidate.display_name,Lead.city.in_(city_values(candidate.city)),Lead.archived.is_(True))))
            if any((lead.facts or {}).get('address',{}).get('value','')==address for lead in archived):raise ServiceError('客户已归档，不能自动重新保存')
        row = db.scalar(select(Lead).where(Lead.dedupe_key == fingerprint))
        if row and row.archived:raise ServiceError('客户已归档，不能自动重新保存')
        if not row:
            possibilities=list(db.scalars(select(Lead).where(Lead.display_name==candidate.display_name,Lead.city.in_(city_values(candidate.city)),Lead.archived.is_(False))))
            if branch:
                # 有明确分店URL的同名商户是不同门店；不能依“无地址”将另一分店电话号码并入原店。
                possibilities=[lead for lead in possibilities if branch in {branch_key(item.get('url','')) for item in lead.evidence or [] if item.get('field')=='name'}]
            compatible=[x for x in possibilities if not address or not (x.facts or {}).get('address',{}).get('value') or (x.facts or {}).get('address',{}).get('value')==address]
            # 仅无地址冲突且候选唯一时补全同一主体，避免缺地址->补地址产生重复；同品牌多店不自动合并。
            if len(compatible)==1:row=compatible[0]
            elif len(possibilities)==1:
                from .evidence_role import same_subject_page
                name_sources={item['source_id']:source_rows[item['source_id']] for item in citations if item['field']=='name'}
                # 同一明确主体页面的地址片段差异保留为争议，不另造一个客户；跨分店目录仍按地址区分。
                if same_subject_page(candidate.display_name,possibilities[0].evidence or [],name_sources):row=possibilities[0]
        if not row:
            row = Lead(dedupe_key=fingerprint, display_name=candidate.display_name, city=candidate.city)
            db.add(row)
            db.flush()
        facts = {**(row.facts or {})}
        for field, value in clean.items():
            if field not in (row.manual_fields or []):
                prior=facts.get(field)
                if prior and prior.get('value')!=value['value'] and not (field=='buyer_role' and row.buyer_role==candidate.buyer_role):
                    # 自动事实冲突只记录双方，不按抓取先后覆盖；人工核验才能消除争议。
                    alternatives=list(prior.get('alternatives',[]))
                    from .contact_view import coexisting_phones
                    parallel=field=='phone' and prior.get('verification') not in {'disputed','inactive'} and coexisting_phones(candidate.display_name,prior,value,source_rows)
                    notice_parallel=bool(field=='procurement' and candidate.procurement_signal
                        and candidate.procurement_signal.source_id==value['source_id'] and prior.get('source_id')!=value['source_id']
                        and prior.get('verification') not in {'disputed','inactive'})
                    if notice_parallel:value={**value,'relation':'additional_notice'};parallel=True
                    if parallel and not notice_parallel:value={**value,'relation':'additional_contact'}
                    if value not in alternatives:alternatives.append(value)
                    facts[field]={**prior,'verification':prior['verification'] if parallel else 'disputed','alternatives':alternatives}
                    if not parallel:
                        from .db import Knowledge
                        for knowledge in db.scalars(select(Knowledge).where(Knowledge.entity_id==row.id)):
                            if prior.get('value','') in knowledge.body or value['value'] in knowledge.body:
                                knowledge.status='disputed';knowledge.updated_at=now()
                else:facts[field]=value if not prior or prior.get('verification')!='disputed' else prior
        row.facts = facts
        row.evidence = list(row.evidence or []) + [x for x in citations if x not in (row.evidence or [])]
        row.legal_name = facts.get('legal_name', {}).get('value')
        if 'industry' not in (row.manual_fields or []):
            # 主营产品原文不是行业名称，缺少行业证据保持空，不把产品标题填成行业。
            row.industry = facts.get('industry', {}).get('value', '')
        row.match_reason, row.matched_products = candidate.match_reason, candidate.matched_products
        if 'buyer_role' not in (row.manual_fields or []):
            if facts.get('buyer_role',{}).get('verification')=='disputed':row.buyer_role='unknown'
            else:row.buyer_role=candidate.buyer_role
        from .role_policy import manufacturer_target
        task=db.get(Job,job_id)
        maker_target=bool(task and manufacturer_target(task.prompt))
        eligible=row.buyer_role in {'dealer','retailer','mall','project_buyer'} or (row.buyer_role=='manufacturer' and maker_target)
        # 角色资格先于资料完整度；仅用户明确厂家画像可匹配制造企业，仍不证明采购；其他任务保持排除。
        verified={k:v for k,v in facts.items() if v.get('verification') not in {'disputed','inactive'}}
        from .contact_view import contact_view
        coverage=10*bool(contact_view(verified,[])['usable_fields'])+10*bool(verified.get('address'))+10*bool(verified.get('categories'))
        product=20*bool(candidate.matched_products)
        signal=0
        from .procurement import signals,signal_score
        prior=dict(row.qualification or {})
        history=signals(prior)
        purchase=candidate.procurement_signal
        if purchase:
            src=source_rows.get(purchase.source_id)
            buyer_quote=purchase.buyer_quote or purchase.quote
            deadline_quote=purchase.deadline_quote or purchase.quote
            excerpts=list(dict.fromkeys([purchase.quote,buyer_quote,*([deadline_quote] if purchase.deadline else [])]))
            if not src or any(re.sub(r'\s+','',quote) not in re.sub(r'\s+','',src['body']) for quote in excerpts) or candidate.display_name not in buyer_quote or purchase.buyer!=candidate.display_name:
                raise ServiceError('采购信号未关联实际客户买方')
            from .procurement import deadline_in_quote
            if purchase.deadline and not deadline_in_quote(purchase.deadline,deadline_quote):raise ServiceError('采购截止时间未匹配原文日期与时间')
            from .knowledge import invalidated_procurement
            if invalidated_procurement(candidate.display_name,purchase,src):
                raise ServiceError('采购公告已被用户标记有误或失效，不能复用旧证据')
            from .procurement import open_buyer_signal
            if open_buyer_signal(candidate.display_name,purchase,src,now()) and ('procurement' not in facts or 'procurement' in verified):signal=10
            # 信号的公告可以独立于基础资料来源，但仍保留原始引文与URL供CRM审计。
            for excerpt in excerpts:
                citation={'field':'procurement','value':excerpt,'quote':excerpt,'source_id':purchase.source_id,
                    'url':src['url'],'title':src['title'],'captured_at':src.get('captured_at',''),'kind':src.get('kind',''),'channel':src.get('channel','')}
                if citation not in row.evidence:row.evidence=[*row.evidence,citation]
            record={**purchase.model_dump(),'intent_score':signal,'url':src['url'],'title':src['title'],'captured_at':src.get('captured_at','')}
            if record not in history:history.append(record)
        primary=purchase.model_dump() if purchase else prior.get('procurement_signal')
        row.qualification={'eligible':eligible,'role':row.buyer_role,'role_score':40 if eligible else 0,
                           'manufacturer_target':maker_target,
                           'product_score':product,'evidence_score':coverage,'signal_score':signal,
                           'procurement_signal':primary,'procurement_signals':history}
        signal=signal_score(row.qualification,facts,now())
        row.qualification={**row.qualification,'signal_score':signal}
        row.score=min(100,40+coverage+product+signal) if eligible else min(35,coverage+product)
        row.updated_at = now()
        job = db.get(Job, job_id)
        if job and row.id not in (job.lead_ids or []):
            job.lead_ids = [*(job.lead_ids or []), row.id]
        return {'id': row.id, 'display_name': row.display_name, 'status': row.status,
                'facts_saved': len(clean), 'score': row.score, 'eligible':eligible}


def refresh_signals(db, lead_id: str | None=None) -> None:
    """db为当前事务，lead_id空检查所有非归档客户；刷新派生采购分，不改原文/公告阶段/业务更新时间。"""
    from .procurement import signal_score
    query=select(Lead).where(Lead.archived.is_(False),Lead.qualification['signal_score'].as_integer()>0)
    if lead_id:query=query.where(Lead.id==lead_id)
    for row in db.scalars(query):
        old=(row.qualification or {}).get('signal_score',0)
        current=signal_score(row.qualification or {},row.facts or {},now())
        if old!=current:
            accounted=old if (row.qualification or {}).get('eligible') else 0
            row.qualification={**row.qualification,'signal_score':current}
            row.score=max(0,(row.score or 0)-accounted+current)
    db.flush()


def lead_view(row) -> dict:
    """row为非空客户ORM行；投影原字段与带时效标签的公告，不推断未知采购事实。"""
    from .procurement import procurement_views
    from .contact_view import contact_view
    records=procurement_views(row.qualification,now())
    for item in records:
        source=next((x for x in row.evidence or [] if x.get('source_id')==item.get('source_id')),{})
        for key in ['url','title','captured_at']:
            if not item.get(key):item[key]=source.get(key,'')
    summary=contact_view(row.facts,row.evidence)
    facts=dict(row.facts or {})
    if facts.get('phone'):
        from .contact_view import phone_info
        phone=facts['phone'];allowed=phone.get('verification') not in {'disputed','inactive'}
        alternatives=[{**item,'contact_href':phone_info(item.get('value'))['href'] if allowed and item.get('relation')=='additional_contact' and item.get('verification') not in {'disputed','inactive'} else ''} for item in phone.get('alternatives',[])]
        facts['phone']={**phone,'contact_href':summary['phone_href'],'alternatives':alternatives}
    if facts.get('wechat'):facts['wechat']={**facts['wechat'],'copyable':'wechat' in summary['usable_fields']}
    return {**serialize(row),'facts':facts,'procurement_signals':records,'contact_summary':summary}


def lead_detail(lead_id: str, conversation_id: str = '') -> dict:
    """读取非空lead_id的详情及跟进；不存在或已归档抛404，返回客户与活动。"""
    with Session.begin() as db:
        refresh_signals(db,lead_id)
        row = db.get(Lead, lead_id)
        if not row or row.archived:
            raise ServiceError('客户线索不存在', 404)
        from .db import Feedback
        result={**lead_view(row), 'feedback': [serialize(x) for x in db.scalars(select(Feedback).where(Feedback.lead_id==lead_id).order_by(Feedback.created_at.desc()))], 'followups': [serialize(x) for x in db.scalars(
            select(Followup).where(Followup.lead_id == lead_id).order_by(Followup.id.desc()))]}
        from .acquisition import acquisition_detail
        result['acquisition']=acquisition_detail(db,result,conversation_id=conversation_id)
        from .lead_analysis import latest_analysis
        result['analysis']=latest_analysis(db,result,conversation_id=result['acquisition']['research_context']['conversation_id'])
        return result


def list_leads(q: str = '', status: str = '', city: str = '', page: int = 1, page_size: int = 20, buyer_role: str='', conversation_id: str='') -> dict:
    """分页查询；空条件不限，page>=1、page_size<=100，返回总数、当前页和全局统计。"""
    with Session.begin() as db:
        refresh_signals(db)
        condition = [Lead.archived.is_(False)]
        if conversation_id:
            linked=func.json_each(Job.lead_ids).table_valued('value')
            condition.append(Lead.id.in_(select(linked.c.value).select_from(Job).join(linked, true()).where(Job.conversation_id==conversation_id)))
        scope_condition=list(condition)
        if q:
            condition.append(Lead.display_name.contains(q, autoescape=True))
        if status:
            condition.append(Lead.status == status)
        if city:
            # 与身份比较保持一致，筛深圳也包含原文深圳市，显示原值不变。
            condition.append(Lead.city.in_(city_values(city)))
        if buyer_role:
            condition.append(Lead.buyer_role==buyer_role)
        total = db.scalar(select(func.count()).select_from(Lead).where(*condition))
        rows = db.scalars(select(Lead).where(*condition).order_by(Lead.qualification['eligible'].as_boolean().is_(True).desc(),func.contact_available(Lead.facts).desc(),Lead.score.desc(), Lead.updated_at.desc())
                          .offset((page - 1) * page_size).limit(page_size))
        stats = {x: db.scalar(select(func.count()).select_from(Lead).where(*scope_condition, Lead.status == x))
                 for x in ['new', 'review', 'contacted', 'nurture', 'converted', 'invalid']}
        items=[lead_view(x) for x in rows]
        from .acquisition import acquisition_views
        acquisition=acquisition_views(db,items)
        return {'items': [{**item,'acquisition':acquisition[item['id']]} for item in items], 'total': total, 'page': page, 'page_size': page_size, 'stats': stats}


def edit_lead(lead_id: str | None, payload: LeadInput) -> dict:
    """手动创建/更新；lead_id空则新建且须名称，payload非空，字段标记人工输入、时间格式校验。"""
    from datetime import datetime
    values = payload.model_dump(exclude_none=True)
    if 'display_name' in values:
        values['display_name'] = values['display_name'].strip()
        if len(values['display_name']) < 2:
            raise ServiceError('客户名称至少需要两个非空字符')
    if values.get('next_followup'):
        try:
            datetime.strptime(values['next_followup'], '%Y-%m-%d %H:%M:%S')
        except ValueError:
            raise ServiceError('下次跟进时间格式应为yyyy-MM-dd HH:mm:ss')
    with Session.begin() as db:
        row = db.get(Lead, lead_id) if lead_id else Lead(display_name=values.get('display_name', ''), dedupe_key=hashlib.sha256(now().encode() + __import__('uuid').uuid4().bytes).hexdigest(), status='new')
        if not row or row.archived:
            raise ServiceError('客户线索不存在', 404)
        if not lead_id and not values.get('display_name'):
            raise ServiceError('新线索需要名称')
        if not lead_id:
            db.add(row)
        facts, manual = {**(row.facts or {})}, set(row.manual_fields or [])
        for key, value in values.items():
            if key in {'phone', 'email', 'wechat', 'address','website','contact_name','contact_position'}:
                facts[key] = {'value': value, 'verification': '人工录入，未独立核验', 'source_id': None}
                manual.add(key)
            else:
                setattr(row, key, value)
                if key in {'industry', 'display_name', 'city', 'buyer_role'}:
                    manual.add('name' if key == 'display_name' else key)
        row.facts, row.manual_fields, row.updated_at = facts, sorted(manual), now()
        if {'buyer_role','phone','email','wechat','address'} & values.keys():
            from .db import Feedback
            row.buyer_role=row.buyer_role or 'unknown'
            blocked=bool(db.scalar(select(Feedback.id).where(Feedback.lead_id==row.id,Feedback.kind=='nonbuyer')))
            eligible=(row.buyer_role in {'dealer','retailer','mall','project_buyer'} or
                (row.buyer_role=='manufacturer' and (row.qualification or {}).get('manufacturer_target',False))) and not blocked
            verified={key:value for key,value in facts.items() if value.get('value') and value.get('verification') not in {'disputed','inactive'}}
            from .contact_view import contact_view
            coverage=10*bool(contact_view(verified,[])['usable_fields'])+10*bool(verified.get('address'))+10*bool(verified.get('categories'))
            product=20*bool(row.matched_products)
            qualification=dict(row.qualification or {})
            purchase=qualification.get('procurement_signal') or {}
            from .procurement import signal_score
            signal=signal_score({**qualification,'eligible':eligible},facts,now())
            qualification.update(eligible=eligible,role=row.buyer_role,role_score=40 if eligible else 0,
                product_score=product,evidence_score=coverage,signal_score=signal)
            if blocked:qualification['reason']='用户标记非买家'
            row.qualification=qualification
            row.score=0 if blocked else min(100,40+coverage+product+signal) if eligible else min(35,coverage+product)
        db.flush()
        return serialize(row)


def import_sample() -> dict:
    """显式导入已研究的深圳灯具历史样本；按网点去重，标资料日期，不宣称实时搜索。"""
    dataset = json.loads((ROOT / 'data/normalized/shenzhen_lighting_20260930.json').read_text(encoding='utf-8'))
    sources = {x['id']: x for x in dataset['sources']}
    imported = 0
    with Session.begin() as db:
        for entity in dataset['entities']:
            key = 'sample:' + entity['id']
            if db.scalar(select(Lead).where(Lead.dedupe_key == key)):
                continue
            facts, quotes = {}, []
            for fact in dataset['facts']:
                if fact['entity_id'] != entity['id'] or fact.get('value') is None:
                    continue
                facts[fact['field']] = {'value': fact['value'], 'verification': '2026-09-30历史研究样本', 'scope': fact.get('qualifier', '')}
                for e in fact.get('evidence', []):
                    src = sources[e['source_id']]
                    quotes.append({**e, 'field': fact['field'], 'value': fact['value'], 'url': src['url'], 'title': src['title'], 'captured_at': '2026-09-30', 'kind': src['kind']})
            db.add(Lead(dedupe_key=key, display_name=entity['name'], city='深圳',
                        industry=facts.get('categories', {}).get('value', ''), facts=facts, evidence=quotes,
                        score=max(20, 90 - entity['priority'] * 2), match_reason=entity['reason'], tags=['历史研究样本', entity['tier']]))
            imported += 1
    return {'imported': imported, 'message': '已导入2026-09-30研究样本，非本次实时搜索'}


def export_workbook(q: str = '', status: str = '', city: str = '', buyer_role: str = ''):
    """导出筛选q/status/city/buyer_role的最多10000条线索；每项空值不限，返回内存xlsx流。

    保留完整事实、口径与逐字段证据；文本公式前加单引号，防止Excel执行外部内容。
    """
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook()
    leads = book.active
    leads.title = '客户线索'
    facts = book.create_sheet('事实与口径')
    sources = book.create_sheet('逐字段证据')
    leads.append(['客户ID', '名称', '城市', '行业', '状态', '负责人', '公开业务电话', '地址', '优先级', '匹配理由', '匹配产品', '主营品类', '经营品牌', '人员规模', '年营收', '采购信息', '供货公司', '采购成交价格', '采购频次', '来源URL', '下次跟进', '更新时间','公开邮箱','联系网站','联系人','联系人职位','信用代码','登记状态','联系信息状态','联系用途','缺口字段','来源平台','数据状态','买家角色','公开业务微信'])
    facts.append(['客户ID','名称','字段','值','核验状态','主体口径','统计日期','来源ID','记录类型','公告阶段','截止时间','参与角色','时效状态','采购人引文','截止时间引文'])
    sources.append(['客户ID','名称','字段','引文','URL','来源标题','采集时间','来源类型','渠道','来源平台'])
    def append(sheet, values):
        """sheet为当前工作表，values为允许导出字段；保护公式文本，无返回值。"""
        sheet.append(["'" + x if isinstance(x,str) and x.startswith(('=','+','-','@')) else x for x in values])
    for page in range(1, 101):
        result = list_leads(q, status, city, page, 100, buyer_role)
        for row in result['items']:
            values = row['facts']
            announcements=row.get('procurement_signals',[])
            procurement=values.get('procurement',{}).get('value','') or '\n'.join(
                f"{item['display_status']} · {item.get('deadline','')}\n{item.get('quote','')}" for item in announcements)
            labels={'phone':'电话','email':'邮箱','website':'联系网站','wechat':'公开业务微信','contact_name':'联系人','contact_position':'职位','credit_code':'信用代码','registry_status':'登记状态','address':'地址','categories':'主营品类','brands':'品牌','headcount':'人员规模','revenue':'年营收','procurement':'采购信息','supplier':'供货公司','purchase_price':'成交价格','frequency':'采购频次','buyer_role':'买家角色'}
            states='；'.join(labels.get(field,field)+'：'+('争议' if fact.get('verification')=='disputed' else '失效') for field,fact in values.items() if fact.get('verification') in {'disputed','inactive'})
            append(leads, [row['id'],row['display_name'],row['city'],row['industry'],row['status'],row['owner'],
                values.get('phone',{}).get('value',''),values.get('address',{}).get('value',''),row['score'],row['match_reason'],
                '、'.join(row['matched_products']),*[procurement if f=='procurement' else values.get(f,{}).get('value','') for f in ['categories','brands','headcount','revenue','procurement','supplier','purchase_price','frequency']],
                '\n'.join(sorted({x['url'] for x in row['evidence'] if x.get('url')})),row['next_followup'],row['updated_at'],
                *[values.get(field,{}).get('value','') for field in ['email','website','contact_name','contact_position','credit_code','registry_status']],
                row['contact_summary']['status'],row['contact_summary']['purpose'],'、'.join(labels.get(x,x) for x in row['contact_summary']['missing_fields']),
                '、'.join(sorted({x['platform'] for x in row['contact_summary']['sources']})),states,row['buyer_role'],values.get('wechat',{}).get('value','')])
            for field, fact in values.items():
                append(facts,[row['id'],row['display_name'],field,fact.get('value',''),fact.get('verification',''),fact.get('scope',''),fact.get('effective_date',''),fact.get('source_id',''),'主值'])
                for alternative in fact.get('alternatives',[]):
                    append(facts,[row['id'],row['display_name'],field,alternative.get('value',''),alternative.get('verification','') if alternative.get('relation')=='additional_contact' else 'disputed · '+alternative.get('verification',''),alternative.get('scope',''),alternative.get('effective_date',''),alternative.get('source_id',''),'其他公开号码' if alternative.get('relation')=='additional_contact' else '替代值'])
            for item in announcements:
                append(facts,[row['id'],row['display_name'],'procurement_signal',item.get('quote',''),item.get('verification','原文匹配'),item.get('buyer',''),'',item.get('source_id',''),'采购公告',item.get('stage',''),item.get('deadline',''),item.get('participant_role',''),item['display_status'],item.get('buyer_quote',''),item.get('deadline_quote','')])
            for e in row['evidence']:
                from .source_catalog import source_label
                append(sources,[row['id'],row['display_name'],e.get('field',''),e.get('quote',e.get('snippet','')),e.get('url',''),e.get('title',''),e.get('captured_at',''),e.get('kind',''),e.get('channel',''),source_label(e.get('url',''))])
        if len(result['items']) < 100: break
    for sheet in book:
        sheet.freeze_panes = 'C2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name='Microsoft YaHei',color='FFFFFF',bold=True)
            cell.fill = PatternFill('solid',fgColor='333333')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 28
            for cell in col[1:]:
                cell.font=Font(name='Microsoft YaHei',size=11)
                cell.alignment = Alignment(vertical='top',wrap_text=True)
    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def tick_clock(job_id: str | None=None) -> None:
    """job_id空更新所有运行任务，否则指定任务；持久累计执行毫秒，停止时不计入等待时间。"""
    import time
    with Session.begin() as db:
        query=select(Job).where(Job.status.in_(['running','cancelling']))
        if job_id:query=query.where(Job.id==job_id)
        stamp=time.time()
        for row in db.scalars(query):
            if row.clock_at:
                row.elapsed_ms=(row.elapsed_ms or 0)+max(0,round((stamp-row.clock_at)*1000))
                row.clock_at=stamp


def settle_interrupted_clock(row, db) -> None:
    """row为遗留运行任务，db为当前事务；只计到最后可证实的活动，不把服务停机时间算进去。"""
    from datetime import datetime, timezone, timedelta
    from .db import ToolCall
    if not row.clock_at:return
    event=db.scalar(select(func.max(JobEvent.created_at)).where(JobEvent.job_id==row.id))
    tool=db.scalar(select(func.max(ToolCall.updated_at)).where(ToolCall.job_id==row.id))
    stamps=[datetime.strptime(x,'%Y-%m-%d %H:%M:%S').replace(tzinfo=timezone(timedelta(hours=8))).timestamp() for x in [event,tool,row.updated_at] if x]
    if stamps:row.elapsed_ms=(row.elapsed_ms or 0)+max(0,round((max(stamps)-row.clock_at)*1000))
    # 原生调用可把心跳一起阻塞，尾段不一定只有5秒；标明这是已确认耗时的下界。
    row.metrics={**(row.metrics or {}),'clock_is_partial':True}
    row.clock_at=None


def job_detail(job_id: str) -> dict:
    """job_id非空，返回最新100条事件、工具轨迹及真实累计时间，不存在404。"""
    import time
    from .db import ToolCall
    with Session() as db:
        row=db.get(Job,job_id)
        if not row:raise ServiceError('任务不存在',404)
        value=serialize(row)
        if row.clock_at and row.status in {'running','cancelling'}:value['elapsed_ms']+=(max(0,round((time.time()-row.clock_at)*1000)))
        events=list(db.scalars(select(JobEvent).where(JobEvent.job_id==job_id).order_by(JobEvent.id.desc()).limit(100)))
        calls=list(db.scalars(select(ToolCall).where(ToolCall.job_id==job_id).order_by(ToolCall.created_at)))
        value['events']=[serialize(x) for x in reversed(events)]
        from .tool_views import tool_view
        value['tools']=[tool_view(x) for x in calls if x.channel != 'model']
        value['source_count']=db.scalar(select(func.count(func.distinct(Evidence.url))).where((Evidence.job_id==job_id)|(Evidence.id.in_((row.metrics or {}).get('source_ids',[])))))
        from .company_models import ResearchDataset
        from .research_dataset import dataset_summary
        value['research_summary']=dataset_summary(list(db.scalars(select(ResearchDataset).where(ResearchDataset.job_id==job_id).order_by(ResearchDataset.created_at,ResearchDataset.id))))
        return value


def job_events(job_id: str, after: int) -> list[dict]:
    """job_id非空，after非负事件ID；每批最多150条按序返回，SSE游标继续，不丢失历史delta。"""
    with Session() as db:
        return [serialize(x) for x in db.scalars(select(JobEvent).where(JobEvent.job_id==job_id,JobEvent.id>after).order_by(JobEvent.id).limit(150))]
